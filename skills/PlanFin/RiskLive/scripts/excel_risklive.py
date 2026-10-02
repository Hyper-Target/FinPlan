"""Excel de RiskLive con el estándar visual y de modelo de PlanFin (ver _shared/convenciones.md).

Hojas: Resumen, Controles, Parametros, Supuestos_EnVivo, Calibracion, Modelo_Base, Simulacion, Resultados,
Percentiles_Mes, Fuentes.

* Modelo_Base es el escenario esperado con FÓRMULAS VIVAS y la estructura del libro de caja del Model Risk del profe
  (Period, Date, Inflow, Outflow, NetCF, InitialBalance, Interest, Balance) más el detalle por activo.
* Resultados, Percentiles_Mes y la hoja Simulacion son valores calculados por Python sobre NSim corridas; la hoja
  Controles avisa si el usuario edita los parámetros y esos resultados dejan de corresponder.
"""
from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "_shared" / "scripts"))
import excel_profe as xp  # noqa: E402

MAX_MESES = 360  # filas de Modelo_Base (30 años)
MUESTRA = 1000
TEAL, AMBAR, ROJO = "0D9488", "F59E0B", "DC2626"
ADVERTENCIA = "Educación financiera, no asesoría de inversión."

COLS_MODELO = ["Period", "Date", "Act", "InflMes", "CumInfl", "TasaCDT", "RetCDT", "RetCOLCAP", "RetUSD", "Rebal",
               "Outflow", "Inflow", "NetCF", "InitialBalance", "Interest", "Balance",
               "IniCDT", "IniCOLCAP", "IniUSD", "IntCDT", "IntCOLCAP", "IntUSD", "BalCDT", "BalCOLCAP", "BalUSD",
               "TotalPreLiq", "RealBalance", "fd", "Xfd"]
COL0_MODELO = 2  # columna B
FILA_HDR_MODELO = 26


def _letra(nombre: str) -> str:
    return get_column_letter(COL0_MODELO + COLS_MODELO.index(nombre))


def filas_modelo() -> list[list]:
    """Fórmulas de Modelo_Base: 361 filas (periodo 0 a 360). Las filas pasado el horizonte quedan en cero."""
    TM = "TModelo[[#This Row],[{}]]".format
    filas = []
    for p in range(0, MAX_MESES + 1):
        r = FILA_HDR_MODELO + 1 + p  # fila de Excel
        rp = r - 1

        def prev(nombre: str) -> str:
            return f"{_letra(nombre)}{rp}"

        act = TM("Act")
        if p == 0:
            fila = {
                "Period": 0,
                "Date": f"=EDATE(FechaInicio,{TM('Period')})",
                "Act": f"=IF(AND({TM('Period')}>=1,{TM('Period')}<=HorizonteMeses),1,0)",
                "InflMes": 0, "CumInfl": 1, "TasaCDT": 0, "RetCDT": 0, "RetCOLCAP": 0, "RetUSD": 0, "Rebal": 0,
                "Outflow": "=AhorroInicial",
                "InitialBalance": 0, "Interest": 0,
                "IniCDT": 0, "IniCOLCAP": 0, "IniUSD": 0, "IntCDT": 0, "IntCOLCAP": 0, "IntUSD": 0,
                "BalCDT": "=PesoCDT*AhorroInicial", "BalCOLCAP": "=PesoCOLCAP*AhorroInicial",
                "BalUSD": "=PesoUSD*AhorroInicial",
            }
        else:
            def ini(peso: str, bal: str) -> str:
                tot = f"({prev('BalCDT')}+{prev('BalCOLCAP')}+{prev('BalUSD')})"
                return f"={act}*IF({TM('Rebal')}=1,{peso}*{tot},{prev(bal)})"

            fila = {
                "Period": p,
                "Date": f"=EDATE(FechaInicio,{TM('Period')})",
                "Act": f"=IF(AND({TM('Period')}>=1,{TM('Period')}<=HorizonteMeses),1,0)",
                "InflMes": f"={act}*(INFb+(INFpi0-INFb)*(1-INFa)^{TM('Period')})",
                "CumInfl": f"=IF({TM('Period')}<=HorizonteMeses,{prev('CumInfl')}*(1+{TM('InflMes')}),0)",
                "TasaCDT": f"={act}*(CDTb+(CDTx0-CDTb)*(1-CDTa)^{TM('Period')}+CDTspread)",
                "RetCDT": f"={act}*((1+{TM('TasaCDT')})^(1/12)-1)*(1-Retencion)",
                "RetCOLCAP": f"={act}*(EXP(MuCOLCAP+SigmaCOLCAP^2/2)-1)",
                "RetUSD": f"={act}*(EXP(MuTRM+SigmaTRM^2/2)-1)",
                "Rebal": f'={act}*IF(Rebalanceo="Mensual",1,IF(Rebalanceo="Anual",IF(MOD({TM("Period")}-1,12)=0,1,0),0))',
                "Outflow": f'={act}*AporteMensual*IF(CrecimientoAporte="Con IPC",{prev("CumInfl")},1)',
                "InitialBalance": f"={prev('Balance')}",
                "Interest": f"={TM('IntCDT')}+{TM('IntCOLCAP')}+{TM('IntUSD')}",
                "IniCDT": ini("PesoCDT", "BalCDT"), "IniCOLCAP": ini("PesoCOLCAP", "BalCOLCAP"),
                "IniUSD": ini("PesoUSD", "BalUSD"),
                "IntCDT": f"={TM('IniCDT')}*{TM('RetCDT')}", "IntCOLCAP": f"={TM('IniCOLCAP')}*{TM('RetCOLCAP')}",
                "IntUSD": f"={TM('IniUSD')}*{TM('RetUSD')}",
                "BalCDT": f"={act}*({TM('IniCDT')}+{TM('IntCDT')}+PesoCDT*{TM('Outflow')})",
                "BalCOLCAP": f"={act}*({TM('IniCOLCAP')}+{TM('IntCOLCAP')}+PesoCOLCAP*{TM('Outflow')})",
                "BalUSD": f"={act}*({TM('IniUSD')}+{TM('IntUSD')}+PesoUSD*{TM('Outflow')})",
            }
        fila["TotalPreLiq"] = f"={TM('BalCDT')}+{TM('BalCOLCAP')}+{TM('BalUSD')}"
        fila["Inflow"] = f"=IF({TM('Period')}=HorizonteMeses,{TM('TotalPreLiq')},0)"
        fila["NetCF"] = f"={TM('Inflow')}-{TM('Outflow')}"
        fila["Balance"] = f"={TM('InitialBalance')}+{TM('Interest')}+{TM('Outflow')}-{TM('Inflow')}"
        fila["RealBalance"] = f"=IF({TM('CumInfl')}>0,{TM('TotalPreLiq')}/{TM('CumInfl')},0)"
        fila["fd"] = f"=(1+PerRate)^(-{TM('Period')})"
        fila["Xfd"] = f"={TM('NetCF')}*{TM('fd')}"
        filas.append([fila[c] for c in COLS_MODELO])
    return filas


def _grafico_limpio(chart) -> None:
    """Cuadrícula gris muy tenue, sin títulos de ejes (que se montan sobre la leyenda) y etiquetas abajo."""
    chart.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="E5E7EB", w=6350))
    chart.x_axis.title = None
    chart.y_axis.title = None
    chart.x_axis.tickLblPos = "low"


def _fila_val(ws, fila, etiqueta, valor, fmt=None, nombre=None, wb=None, unidad="", obs=""):
    xp.fila_parametro(ws, fila, etiqueta, valor, unidad, obs, nombre, wb, fmt)


def construir_excel(ctx: dict, ruta: Path) -> None:
    p, cal, st = ctx["params"], ctx["cal"], ctx["stats"]
    wb = Workbook()
    ws_res = wb.active
    ws_res.title = "Resumen"
    ws_ctl = wb.create_sheet("Controles")
    ws_par = wb.create_sheet("Parametros")
    ws_sup = wb.create_sheet("Supuestos_EnVivo")
    ws_cal = wb.create_sheet("Calibracion")
    ws_mod = wb.create_sheet("Modelo_Base")
    ws_sim = wb.create_sheet("Simulacion")
    ws_rsl = wb.create_sheet("Resultados")
    ws_pct = wb.create_sheet("Percentiles_Mes")
    ws_fue = wb.create_sheet("Fuentes")
    wb.calculation.fullCalcOnLoad = True
    meses = p["meses"]

    # ------------------------------------------------------------------ Parametros
    xp.titulo(ws_par, "B2", "Parámetros (celdas azules: edítelas y Modelo_Base se recalcula)", 14)
    for j, h in enumerate(["Variable", "Valor", "Unidad", "Observación"]):
        xp.encabezado(ws_par.cell(row=4, column=2 + j, value=h))
    P = xp.fila_parametro
    P(ws_par, 5, "Ahorro inicial", p["ahorro0"], "COP", "Depósito en el periodo 0", "AhorroInicial", wb, xp.FMT_COP, True)
    P(ws_par, 6, "Aporte mensual", p["aporte"], "COP", "Al final de cada mes", "AporteMensual", wb, xp.FMT_COP, True)
    P(ws_par, 7, "Horizonte", p["horizonte"], "años", f"Hasta {MAX_MESES // 12} años", "HorizonteAnios", wb, "0", True)
    P(ws_par, 8, "Horizonte en meses", "=HorizonteAnios*12", "meses", "Calculado", "HorizonteMeses", wb, "0")
    P(ws_par, 9, "Meta", p["meta"], "COP de hoy", "Valor real que se quiere alcanzar", "Meta", wb, xp.FMT_COP, True)
    P(ws_par, 10, "Peso en CDT", p["pesos"][0], "%", "Renta fija", "PesoCDT", wb, xp.FMT_PCT, True)
    P(ws_par, 11, "Peso en COLCAP", p["pesos"][1], "%", "Renta variable local", "PesoCOLCAP", wb, xp.FMT_PCT, True)
    P(ws_par, 12, "Peso en dólares", p["pesos"][2], "%", "Efectivo en USD, rinde la variación de la TRM", "PesoUSD", wb,
      xp.FMT_PCT, True)
    P(ws_par, 13, "Suma de pesos", "=PesoCDT+PesoCOLCAP+PesoUSD", "%", "Debe ser 100 %", "SumaPesos", wb, xp.FMT_PCT)
    P(ws_par, 14, "Rebalanceo", p["rebalanceo"].capitalize(), "", "Mensual, Anual o Ninguno", "Rebalanceo", wb, None, True)
    P(ws_par, 15, "Crecimiento del aporte", "Con IPC" if p["con_ipc"] else "Fijo", "", "Con IPC mantiene el aporte en pesos reales",
      "CrecimientoAporte", wb, None, True)
    P(ws_par, 16, "Retención sobre rendimiento de CDT", p["retencion"], "%", "Verifique la tarifa vigente", "Retencion", wb,
      xp.FMT_PCT, True)
    P(ws_par, 17, "Fecha de inicio", p["fecha_inicio"], "", "Periodo 0", "FechaInicio", wb, xp.FMT_FECHA, True)
    P(ws_par, 19, "Simulaciones (NSim)", p["n_sim"], "corridas", "Solo cambia volviendo a correr el script", "NSim", wb, xp.FMT_COP)
    P(ws_par, 20, "Semilla (Seed)", p["seed"], "", "Misma semilla, mismo resultado", "Seed", wb, "0")
    P(ws_par, 21, "Distribución de los shocks", "t de Student" if p["dist"] == "t" else "Normal", "",
      "Solo cambia volviendo a correr el script", "Distribucion", wb)
    P(ws_par, 22, "Ventana de calibración", p["ventana"], "años", "Solo cambia volviendo a correr el script", "VentanaAnios", wb, "0")
    P(ws_par, 23, "Fecha de los datos", cal["fecha_datos"], "", "Dato diario más antiguo usado", "FechaDatos", wb, xp.FMT_FECHA)
    P(ws_par, 24, "Fecha de consulta", ctx["ahora"].strftime("%Y-%m-%d %H:%M"), "", "", "FechaConsulta", wb)
    for celda, lista in (("C14", '"Mensual,Anual,Ninguno"'), ("C15", '"Con IPC,Fijo"')):
        dv = DataValidation(type="list", formula1=lista, allow_blank=False)
        ws_par.add_data_validation(dv)
        dv.add(celda)
    xp.ajustar_anchos(ws_par)
    ws_par.column_dimensions["B"].width = 34
    ws_par.column_dimensions["E"].width = 52

    # ------------------------------------------------------------------ Supuestos_EnVivo
    xp.titulo(ws_sup, "A2", "Supuestos que usó el agente hoy, con su fuente", 14)
    cols_s = ["Variable", "Valor", "Unidad", "Fuente", "IDSerie", "FechaDato", "FechaConsulta"]
    filas_s = [[s["Variable"], s["Valor"], s["Unidad"], s["Fuente"], s["IDSerie"], s["FechaDato"],
                ctx["ahora"].strftime("%Y-%m-%d %H:%M")] for s in ctx["supuestos"]]
    xp.escribir_tabla(ws_sup, "TSupuestos", cols_s, filas_s, 4, 1, {"FechaDato": xp.FMT_FECHA})
    for i, s in enumerate(ctx["supuestos"]):
        ws_sup.cell(row=5 + i, column=2).number_format = s["Formato"]
        ws_sup.cell(row=5 + i, column=2).alignment = Alignment(horizontal="right")
    xp.ajustar_anchos(ws_sup, 12, 70)
    ws_sup.column_dimensions["A"].width = 40

    # ------------------------------------------------------------------ Calibracion
    xp.titulo(ws_cal, "B2", "Calibración con datos oficiales (ventana de %d años, datos mensuales)" % p["ventana"], 14)
    for j, h in enumerate(["Parámetro", "Valor", "Unidad", "Nota"]):
        xp.encabezado(ws_cal.cell(row=4, column=2 + j, value=h))
    C = cal
    filas_c = [
        ("Media del retorno log mensual COLCAP", C["COLCAP"]["mu"], "0.0000%", "MuCOLCAP", "mensual", "MCO / promedio de la ventana"),
        ("Volatilidad mensual COLCAP", C["COLCAP"]["sigma"], "0.0000%", "SigmaCOLCAP", "mensual", ""),
        ("Media del retorno log mensual TRM", C["TRM"]["mu"], "0.0000%", "MuTRM", "mensual", ""),
        ("Volatilidad mensual TRM", C["TRM"]["sigma"], "0.0000%", "SigmaTRM", "mensual", ""),
        ("CDT: tasa del sistema hoy (x0)", C["CDT"]["x0"], "0.00%", "CDTx0", "% EA", "BanRep serie 240"),
        ("CDT: prima de la mejor entidad (spread)", C["CDT"]["spread"], "0.00%", "CDTspread", "% EA",
         C["CDT"]["origen"]),
        ("CDT: tasa de partida (mejor de hoy)", C["CDT"]["y0"], "0.00%", "CDTy0", "% EA", "x0 + spread"),
        ("CDT: velocidad de reversión (a)", C["CDT"]["a"], "0.0000", "CDTa", "mensual", f"Modo: {C['CDT']['modo']}"),
        ("CDT: nivel de largo plazo (b)", C["CDT"]["b"], "0.00%", "CDTb", "% EA", "Reversión a la media (Vasicek discreto)"),
        ("CDT: volatilidad mensual (sigma)", C["CDT"]["sigma"], "0.0000%", "CDTsigma", "% EA", "Desviación de los residuos MCO"),
        ("Inflación mensual de partida (pi0)", C["IPC"]["pi0"], "0.0000%", "INFpi0", "mensual", "Equivalente mensual del IPC a 12 meses"),
        ("Inflación: velocidad de reversión (a)", C["IPC"]["a"], "0.0000", "INFa", "mensual", ""),
        ("Inflación: nivel de largo plazo (b)", C["IPC"]["b"], "0.0000%", "INFb", "mensual", "Meta de inflación del BanRep en base mensual"),
        ("Inflación: volatilidad mensual (sigma)", C["IPC"]["sigma"], "0.0000%", "INFsigma", "mensual", ""),
        ("Observaciones mensuales mínimas por serie", C["n_obs_min"], "0", "NObsMin", "meses", "Mínimo exigido: 60"),
        ("Grados de libertad t de Student", C["df_t"], "0.00", "DfT", "", "Solo se usan con distribución t"),
    ]
    for i, (et, val, fmt, nombre, un, nota) in enumerate(filas_c):
        _fila_val(ws_cal, 5 + i, et, val, fmt, nombre, wb, un, nota)
    r0 = 5 + len(filas_c) + 2
    ws_cal.cell(row=r0, column=2, value="Pruebas de normalidad de los retornos mensuales").font = Font(bold=True, size=8, color=xp.GRIS_TENUE)
    enc = ["Serie", "Asimetría", "Curtosis (exceso)", "Jarque-Bera", "Valor p", "Normal rechazada (5 %)", "Observaciones"]
    for j, h in enumerate(enc):
        cel = xp.encabezado(ws_cal.cell(row=r0 + 1, column=2 + j, value=h)) or ws_cal.cell(row=r0 + 1, column=2 + j)
        if j > 0:
            cel.alignment = Alignment(horizontal="right", wrap_text=True, vertical="center")
    for i, k in enumerate(("COLCAP", "TRM")):
        m = C[k]
        vals = [k, m["asimetria"], m["curtosis_exceso"], m["jb"], m["jb_p"], "Sí" if m["normal_rechazada"] else "No", m["n_obs"]]
        for j, v in enumerate(vals):
            cel = ws_cal.cell(row=r0 + 2 + i, column=2 + j, value=v)
            cel.border = xp._BORDE
            if isinstance(v, float):
                cel.number_format = "0.000"
    r1 = r0 + 5
    ws_cal.cell(row=r1, column=2, value="Matriz de correlación de los shocks (estandarizados)").font = Font(bold=True, size=8, color=xp.GRIS_TENUE)
    nombres = ["COLCAP", "TRM", "CDT", "IPC"]
    xp.encabezado(ws_cal.cell(row=r1 + 1, column=2, value=""))
    for j, nm in enumerate(nombres):
        cel = ws_cal.cell(row=r1 + 1, column=3 + j, value=nm)
        xp.encabezado(cel)
        cel.alignment = Alignment(horizontal="right")
    for i, nm in enumerate(nombres):
        ws_cal.cell(row=r1 + 2 + i, column=2, value=nm).border = xp._BORDE
        for j in range(4):
            cel = ws_cal.cell(row=r1 + 2 + i, column=3 + j, value=float(C["corr"]["matriz"][i][j]))
            cel.number_format = "0.000"
            cel.border = xp._BORDE
    ws_cal.cell(row=r1 + 7, column=2, value="Definida positiva").font = Font(bold=True)
    ws_cal.cell(row=r1 + 7, column=3, value="Sí" if C["corr"]["psd_ok"] else "No, corregida a la PSD más cercana")
    ws_cal.cell(row=r1 + 8, column=2, value="Autovalor mínimo").font = Font(bold=True)
    ws_cal.cell(row=r1 + 8, column=3, value=float(C["corr"]["min_autovalor"])).number_format = "0.0000"
    xp.definir_nombre(wb, "CorrPSD", f"'Calibracion'!$C${r1 + 7}")
    ws_cal.column_dimensions["B"].width = 44
    for col in "CDEFGH":
        ws_cal.column_dimensions[col].width = 16
    ws_cal.column_dimensions["E"].width = 22
    ws_cal.column_dimensions["F"].width = 40

    # ------------------------------------------------------------------ Modelo_Base
    xp.cabecera_modelo(ws_mod, "Modelo base: escenario esperado con fórmulas vivas",
                       "Libro de caja con la estructura del Model Risk del curso (perspectiva del inversionista); "
                       "los rendimientos son los valores esperados de la calibración.")
    for j, h in ((2, "ITEM"), (3, "DESCRIPTION"), (6, "VARIABLE"), (7, "VALUE"), (8, "UNIT")):
        xp.encabezado(ws_mod.cell(row=6, column=j, value=h))
    for c_ in (4, 5):
        xp.encabezado(ws_mod.cell(row=6, column=c_))
    items = [
        ("Type of asset", None, "Portafolio de ahorro", None, ""),
        ("Value of the asset", "Asset", "=AhorroInicial", xp.FMT_COP, "COP"),
        ("Downpayment (%)", "DownpaymentP", 0, xp.FMT_PCT, ""),
        ("Downpayment ($)", "Downpayment", "=Asset*DownpaymentP", xp.FMT_COP, "COP"),
        ("Value invested (Principal)", "Principal", "=Asset-Downpayment", xp.FMT_COP, "COP"),
        ("Term (years)", "Term", "=HorizonteAnios", "0", "años"),
        ("Frequency of payments", "Freq", "Monthly", None, ""),
        ("Number of payments per year", "m", 12, "0", ""),
        ("Effective annual rate (EAR) implícita", "EAR", "=(1+PerRate)^m-1", xp.FMT_PCT, "% EA"),
        ("Nominal annual rate (NOM) implícita", "NOM", "=m*((1+EAR)^(1/m)-1)", xp.FMT_PCT, "NAMV"),
        ("Periodic rate (TIR mensual del flujo)", "PerRate", "=IRR(TModelo[NetCF],0.01)", "0.0000%", "EM"),
        ("Disbursement date", "Disbursement", "=FechaInicio", xp.FMT_FECHA, ""),
        ("Number of installments", "nInstallments", "=HorizonteMeses", "0", "aportes"),
        ("Value of the installment (aporte inicial)", "Installment", "=AporteMensual", xp.FMT_COP, "COP"),
        ("Perspective", None, "Investor", None, ""),
    ]
    for i, (desc, var, val, fmt, un) in enumerate(items):
        r = 7 + i
        ws_mod.cell(row=r, column=2, value=i + 1).border = xp._BORDE
        ws_mod.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
        ws_mod.cell(row=r, column=3, value=desc)
        ws_mod.cell(row=r, column=6, value=var or "")
        cel = ws_mod.cell(row=r, column=7, value=val)
        cel.alignment = Alignment(horizontal="right")
        if fmt:
            cel.number_format = fmt
        ws_mod.cell(row=r, column=8, value=un)
        for c_ in range(2, 9):
            ws_mod.cell(row=r, column=c_).border = xp._BORDE
        if var:
            xp.definir_nombre(wb, var, f"'Modelo_Base'!$G${r}")
    salidas = [
        ("Saldo final nominal (antes de liquidar)", f"=INDEX(TModelo[TotalPreLiq],HorizonteMeses+1)", xp.FMT_COP, "FinalNominal"),
        ("Saldo final real (pesos de hoy)", f"=INDEX(TModelo[RealBalance],HorizonteMeses+1)", xp.FMT_COP, "FinalReal"),
        ("LastBalance (debe ser 0 tras liquidar)", f"=INDEX(TModelo[Balance],HorizonteMeses+1)", "0.000000", "LastBalance"),
        ("Sum Xfd (debe ser 0)", "=SUM(TModelo[Xfd])", "0.000000", "SumXfd"),
        ("Total aportado (nominal)", "=SUM(TModelo[Outflow])", xp.FMT_COP, "TotalAportado"),
    ]
    ws_mod["J6"] = "SALIDAS DEL ESCENARIO ESPERADO"
    ws_mod["J6"].font = Font(bold=True, size=8, color=xp.GRIS_TENUE)
    for i, (et, f, fmt, nom) in enumerate(salidas):
        r = 7 + i
        ws_mod.merge_cells(start_row=r, start_column=10, end_row=r, end_column=13)
        ws_mod.cell(row=r, column=10, value=et)
        cel = ws_mod.cell(row=r, column=14, value=f)
        cel.number_format = fmt
        cel.alignment = Alignment(horizontal="right")
        for c_ in range(10, 15):
            ws_mod.cell(row=r, column=c_).border = xp._BORDE
        xp.definir_nombre(wb, nom, f"'Modelo_Base'!$N${r}")
    ws_mod["B24"] = "LIBRO DE CAJA · PERIODO 0 = DEPÓSITO INICIAL, APORTES AL FINAL DE CADA MES · FILAS FUERA DEL HORIZONTE EN GRIS"
    ws_mod["B24"].font = Font(bold=True, size=8, color=xp.GRIS_TENUE)
    grupos = [("Period", "CumInfl", "CALENDARIO Y MACRO"),
              ("TasaCDT", "Rebal", "RENDIMIENTOS ESPERADOS"),
              ("Outflow", "Balance", "LIBRO DE CAJA (ESTRUCTURA DEL PROFE)"),
              ("IniCDT", "IntUSD", "INTERESES POR ACTIVO"), ("BalCDT", "BalUSD", "SALDOS POR ACTIVO"),
              ("TotalPreLiq", "Xfd", "SALIDAS")]
    for a, b, txt in grupos:
        ca, cb = COLS_MODELO.index(a) + COL0_MODELO, COLS_MODELO.index(b) + COL0_MODELO
        ws_mod.merge_cells(start_row=FILA_HDR_MODELO - 1, start_column=ca, end_row=FILA_HDR_MODELO - 1, end_column=cb)
        c_ = ws_mod.cell(row=FILA_HDR_MODELO - 1, column=ca, value=txt)
        c_.font = Font(bold=True, size=8, color=xp.ACENTO)
        c_.alignment = Alignment(horizontal="left")
    fmt_m = {"Date": xp.FMT_FECHA, "InflMes": "0.000%", "CumInfl": "0.0000", "TasaCDT": "0.00%", "RetCDT": "0.000%",
             "RetCOLCAP": "0.000%", "RetUSD": "0.000%", "Outflow": xp.FMT_COP, "Inflow": xp.FMT_COP, "NetCF": xp.FMT_COP,
             "InitialBalance": xp.FMT_COP, "Interest": xp.FMT_COP, "Balance": xp.FMT_COP, "IniCDT": xp.FMT_COP,
             "IniCOLCAP": xp.FMT_COP, "IniUSD": xp.FMT_COP, "IntCDT": xp.FMT_COP, "IntCOLCAP": xp.FMT_COP,
             "IntUSD": xp.FMT_COP, "BalCDT": xp.FMT_COP, "BalCOLCAP": xp.FMT_COP, "BalUSD": xp.FMT_COP,
             "TotalPreLiq": xp.FMT_COP, "RealBalance": xp.FMT_COP, "fd": "0.000000", "Xfd": xp.FMT_NUM}
    xp.escribir_tabla(ws_mod, "TModelo", COLS_MODELO, filas_modelo(), FILA_HDR_MODELO, COL0_MODELO, fmt_m)
    ultima = FILA_HDR_MODELO + 1 + MAX_MESES
    ws_mod.conditional_formatting.add(
        f"B{FILA_HDR_MODELO + 2}:{get_column_letter(COL0_MODELO + len(COLS_MODELO) - 1)}{ultima}",
        FormulaRule(formula=[f"$B{FILA_HDR_MODELO + 2}>HorizonteMeses"], font=Font(color="C0C4CC")))
    ws_mod.freeze_panes = ws_mod.cell(row=FILA_HDR_MODELO + 1, column=COL0_MODELO + 2)
    for i in range(len(COLS_MODELO)):
        ws_mod.column_dimensions[get_column_letter(COL0_MODELO + i)].width = 13
    ws_mod.column_dimensions["B"].width = 8
    ws_mod.column_dimensions["C"].width = 12

    # ------------------------------------------------------------------ Simulacion
    xp.titulo(ws_sim, "A2", f"Muestra de las primeras {MUESTRA:,} corridas de {p['n_sim']:,} (las estadísticas usan todas)".replace(",", "."), 14)
    m = ctx["muestra"]
    cols_m = ["Corrida", "ValorFinalNominal", "AportesReales", "ValorFinalReal", "GananciaReal", "CumpleMeta", "MaxDrawdown"]
    TS = "TSim[[#This Row],[{}]]".format
    filas_m = []
    for i in range(len(m["final_real"])):
        filas_m.append([i + 1, float(m["final_nominal"][i]), float(m["aportes_reales"][i]), float(m["final_real"][i]),
                        f"={TS('ValorFinalReal')}-{TS('AportesReales')}",
                        f"=IF({TS('ValorFinalReal')}>=Meta,1,0)", float(m["drawdown"][i])])
    xp.escribir_tabla(ws_sim, "TSim", cols_m, filas_m, 8, 1,
                      {"ValorFinalNominal": xp.FMT_COP, "AportesReales": xp.FMT_COP, "ValorFinalReal": xp.FMT_COP,
                       "GananciaReal": xp.FMT_COP, "CumpleMeta": "0", "MaxDrawdown": "0.00%"})
    ws_sim["A4"] = "ESTADÍSTICAS DE LA MUESTRA (EN VIVO)"
    ws_sim["A4"].font = Font(bold=True, size=8, color=xp.GRIS_TENUE)
    est = [("P(cumplir la meta) en la muestra", "=AVERAGE(TSim[CumpleMeta])", xp.FMT_PCT, "PMetaMuestra"),
           ("Mediana del valor real", "=MEDIAN(TSim[ValorFinalReal])", xp.FMT_COP, "MedianaMuestra"),
           ("Percentil 5 del valor real", "=PERCENTILE(TSim[ValorFinalReal],0.05)", xp.FMT_COP, "P5Muestra")]
    ws_sim["A5"], ws_sim["B5"] = est[0][0], est[0][1]
    ws_sim["A6"], ws_sim["B6"] = est[1][0], est[1][1]
    ws_sim["A7"], ws_sim["B7"] = est[2][0], est[2][1]
    for r_, (_, _, fmt, nom) in zip((5, 6, 7), est):
        ws_sim.cell(row=r_, column=2).number_format = fmt
        xp.definir_nombre(wb, nom, f"'Simulacion'!$B${r_}")
    ws_sim.freeze_panes = "A9"
    xp.ajustar_anchos(ws_sim, 14, 40)
    ws_sim.column_dimensions["A"].width = 34

    # ------------------------------------------------------------------ Resultados
    xp.titulo(ws_rsl, "B2", "Resultados de la simulación (valores calculados por Python en la corrida)", 14)
    for j, h in enumerate(["Resultado", "Valor", "Unidad", "Observación"]):
        xp.encabezado(ws_rsl.cell(row=4, column=2 + j, value=h))
    pc = st["percentiles"]
    res_rows = [
        ("P(cumplir la meta)", st["p_meta"], xp.FMT_PCT, "PMeta", "% de corridas", "Valor real final >= meta"),
        ("Percentil 5 del valor real", pc[5], xp.FMT_COP, "P5Real", "COP de hoy", "Peor caso razonable"),
        ("Percentil 25", pc[25], xp.FMT_COP, "P25Real", "COP de hoy", ""),
        ("Mediana del valor real", pc[50], xp.FMT_COP, "MedianaReal", "COP de hoy", ""),
        ("Percentil 75", pc[75], xp.FMT_COP, "P75Real", "COP de hoy", ""),
        ("Percentil 95", pc[95], xp.FMT_COP, "P95Real", "COP de hoy", ""),
        ("Media del valor real", st["media"], xp.FMT_COP, "MediaReal", "COP de hoy", ""),
        ("Valor final nominal (mediana)", st["final_nominal_mediana"], xp.FMT_COP, "MedianaNominal", "COP futuros", ""),
        ("Aportes reales medios", st["aportes_reales_medio"], xp.FMT_COP, "AportesRealesMedios", "COP de hoy", "Ahorro inicial + aportes en pesos de hoy"),
        ("VaR 95 % de la ganancia real", st["var95"], '#,##0;[Color10]-#,##0', "VaR95", "COP de hoy",
         "Pérdida real en el percentil 95; negativo = aún hay ganancia real en el peor 5 %"),
        ("CVaR 95 % de la ganancia real", st["cvar95"], '#,##0;[Color10]-#,##0', "CVaR95", "COP de hoy", "Promedio de la cola del peor 5 %"),
        ("P(perder poder adquisitivo)", st["p_perdida_real"], xp.FMT_PCT, "PPerdidaReal", "% de corridas", "Valor real final < aportes reales"),
        ("Caída máxima mediana del saldo real", st["dd_mediana"], xp.FMT_PCT, "DDMediana", "%", "Desde el máximo previo"),
        ("Caída máxima, percentil 95", st["dd_p95"], xp.FMT_PCT, "DDP95", "%", ""),
        ("Aporte mensual para 80 % de probabilidad", ctx["aporte_80"], xp.FMT_COP, "Aporte80", "COP por mes", "Misma semilla; solución exacta por cuantil"),
    ]
    for i, (et, val, fmt, nom, un, ob) in enumerate(res_rows):
        _fila_val(ws_rsl, 5 + i, et, val, fmt, nom, wb, un, ob)
    rr = 5 + len(res_rows) + 1
    ws_rsl.cell(row=rr, column=2, value="PARÁMETROS DE LA CORRIDA (Controles avisa si Parametros se edita después)").font = Font(bold=True, size=8, color=xp.GRIS_TENUE)
    corr_rows = [("AhorroInicialCorrida", p["ahorro0"], xp.FMT_COP), ("AporteMensualCorrida", p["aporte"], xp.FMT_COP),
                 ("HorizonteMesesCorrida", meses, "0"), ("MetaCorrida", p["meta"], xp.FMT_COP),
                 ("PesoCDTCorrida", p["pesos"][0], xp.FMT_PCT), ("PesoCOLCAPCorrida", p["pesos"][1], xp.FMT_PCT),
                 ("PesoUSDCorrida", p["pesos"][2], xp.FMT_PCT), ("RebalanceoCorrida", p["rebalanceo"].capitalize(), None),
                 ("CrecimientoCorrida", "Con IPC" if p["con_ipc"] else "Fijo", None), ("RetencionCorrida", p["retencion"], xp.FMT_PCT)]
    for i, (nom, val, fmt) in enumerate(corr_rows):
        _fila_val(ws_rsl, rr + 1 + i, nom, val, fmt, nom, wb)
    rs = rr + len(corr_rows) + 3
    ws_rsl.cell(row=rs - 1, column=2, value="SENSIBILIDAD AL APORTE MENSUAL (misma simulación)").font = Font(bold=True, size=8, color=xp.GRIS_TENUE)
    sa = ctx["sens_aporte"]
    xp.escribir_tabla(ws_rsl, "TSensAporte", ["CambioAporte", "AporteMensual", "PMeta", "MedianaReal"],
                      [[r["cambio"], r["aporte"], r["p_meta"], r["mediana"]] for r in sa], rs, 2,
                      {"CambioAporte": "+0%;-0%;0%", "AporteMensual": xp.FMT_COP, "PMeta": xp.FMT_PCT, "MedianaReal": xp.FMT_COP})
    rc = rs + len(sa) + 3
    ws_rsl.cell(row=rc - 1, column=2, value="SENSIBILIDAD AL PESO EN COLCAP (el resto se reparte en la proporción original)").font = Font(bold=True, size=8, color=xp.GRIS_TENUE)
    sc = ctx["sens_col"]
    xp.escribir_tabla(ws_rsl, "TSensCOLCAP", ["PesoCOLCAP", "PesoCDT", "PesoUSD", "PMeta", "MedianaReal", "P5Real"],
                      [[r["w_col"], r["pesos"][0], r["pesos"][2], r["p_meta"], r["mediana"], r["p5"]] for r in sc], rc, 2,
                      {"PesoCOLCAP": "0%", "PesoCDT": "0%", "PesoUSD": "0%", "PMeta": xp.FMT_PCT, "MedianaReal": xp.FMT_COP, "P5Real": xp.FMT_COP})
    rh = rc + len(sc) + 3
    ws_rsl.cell(row=rh - 1, column=2, value="DISTRIBUCIÓN DEL VALOR REAL FINAL (30 intervalos entre los percentiles 0,5 y 99,5)").font = Font(bold=True, size=8, color=xp.GRIS_TENUE)
    cuentas, bordes = ctx["hist"]
    filas_h = []
    for i, cnt in enumerate(cuentas):
        mid = (bordes[i] + bordes[i + 1]) / 2
        filas_h.append([f"{mid / 1e6:.1f}".replace(".", ","), float(bordes[i]), float(bordes[i + 1]), int(cnt),
                        "Sobre la meta" if mid >= p["meta"] else "Bajo la meta"])
    xp.escribir_tabla(ws_rsl, "THist", ["Centro", "Desde", "Hasta", "Corridas", "Zona"], filas_h, rh, 2,
                      {"Desde": xp.FMT_COP, "Hasta": xp.FMT_COP, "Corridas": xp.FMT_COP})
    ws_rsl.column_dimensions["B"].width = 42
    for col, w in zip("CDEFG", (18, 16, 22, 40, 18)):
        ws_rsl.column_dimensions[col].width = w
    ws_rsl.column_dimensions["F"].width = 24
    ws_rsl.column_dimensions["E"].width = 26

    # ------------------------------------------------------------------ Percentiles_Mes
    xp.titulo(ws_pct, "A2", "Percentiles del saldo real mes a mes (tabla detrás del fan chart)", 14)
    pm = ctx["pct_mes"]
    filas_p = [[t] + [float(x) for x in pm[t]] + ["=MetaCorrida"] for t in range(pm.shape[0])]
    xp.escribir_tabla(ws_pct, "TPercMes", ["Mes", "P5", "P25", "P50", "P75", "P95", "Meta"], filas_p, 4, 1,
                      {c: xp.FMT_COP for c in ("P5", "P25", "P50", "P75", "P95", "Meta")})
    ws_pct.freeze_panes = "A5"
    xp.ajustar_anchos(ws_pct, 14, 24)

    # ------------------------------------------------------------------ Fuentes
    xp.titulo(ws_fue, "A2", "Fuentes, validaciones y notas metodológicas", 14)
    filas_f = [[f["Fuente"], f["URL"], f["FechaHoraConsulta"], f["Observaciones"], f["Detalle"]] for f in ctx["fuentes"]]
    xp.escribir_tabla(ws_fue, "TFuentes", ["Fuente", "URL", "FechaHoraConsulta", "Observaciones", "Detalle"], filas_f, 4, 1,
                      {"Observaciones": xp.FMT_COP})
    r0 = 4 + len(filas_f) + 3
    xp.escribir_tabla(ws_fue, "TValidaciones", ["Validacion", "Resultado", "Detalle"],
                      [[c["Validacion"], c["Resultado"], c["Detalle"]] for c in ctx["checks"]], r0, 1)
    r1 = r0 + len(ctx["checks"]) + 3
    notas = ctx["notas_metodologicas"] + [f"ADVERTENCIA DE CALIBRACIÓN: {a}" for a in cal["avisos"]] + [ADVERTENCIA]
    for i, t in enumerate(notas):
        xp.nota(ws_fue, f"A{r1 + i}", t, 5, 40)
    xp.ajustar_anchos(ws_fue, 14, 60)
    ws_fue.column_dimensions["B"].width = 70

    # ------------------------------------------------------------------ Controles
    xp.titulo(ws_ctl, "B2", "Controles de integridad del modelo", 16)
    ctl = [
        ["C01", "Los pesos del portafolio suman 100 %", '=IF(ABS(SumaPesos-1)<0.0001,"OK","ERROR: los pesos no suman 100 %")'],
        ["C02", "Tras liquidar en el horizonte el saldo del libro de caja es cero (LastBalance)", '=IF(ABS(LastBalance)<1,"OK","ERROR")'],
        ["C03", "Suma de flujos descontados a la TIR es cero (ΣXfd)", '=IF(ABS(SumXfd)<1,"OK","ERROR")'],
        ["C04", "El libro de caja no contiene errores de fórmula", '=IF(SUMPRODUCT(--ISERROR(TModelo[[Period]:[Xfd]]))=0,"OK","ERROR")'],
        ["C05", "El horizonte está entre 1 y 30 años", f'=IF(AND(HorizonteAnios>=1,HorizonteAnios<={MAX_MESES // 12}),"OK","ERROR")'],
        ["C06", "La matriz de correlación es definida positiva", '=IF(CorrPSD="Sí","OK","ADVERTENCIA: se corrigió a la PSD más cercana")'],
        ["C07", "Al menos 60 observaciones mensuales por serie", '=IF(NObsMin>=60,"OK","ERROR")'],
        ["C08", "Dato de mercado con menos de 7 días de antigüedad", '=IF(TODAY()-FechaDatos<7,"OK","ADVERTENCIA: dato con "&(TODAY()-FechaDatos)&" días")'],
        ["C09", "La probabilidad de cumplir la meta está entre 0 y 1", '=IF(AND(PMeta>=0,PMeta<=1),"OK","ERROR")'],
        ["C10", "Los parámetros no se editaron después de la corrida",
         '=IF(AND(AhorroInicial=AhorroInicialCorrida,AporteMensual=AporteMensualCorrida,HorizonteMeses=HorizonteMesesCorrida,'
         'Meta=MetaCorrida,PesoCDT=PesoCDTCorrida,PesoCOLCAP=PesoCOLCAPCorrida,PesoUSD=PesoUSDCorrida,'
         'Rebalanceo=RebalanceoCorrida,CrecimientoAporte=CrecimientoCorrida,Retencion=RetencionCorrida),"OK",'
         '"ADVERTENCIA: parámetros editados; Resultados corresponde a la corrida original")'],
        ["C11", "La muestra de 1.000 corridas es consistente con la simulación completa",
         '=IF(Meta<>MetaCorrida,"ADVERTENCIA: la Meta cambió",IF(ABS(PMetaMuestra-PMeta)<=4*SQRT(MAX(PMeta*(1-PMeta),0.0001)/COUNT(TSim[ValorFinalReal])),"OK","ADVERTENCIA: la muestra se aleja de la simulación completa"))'],
        ["C12", "Conciliación del Excel contra un cálculo independiente en Python", ctx.get("estado_conciliacion", "OMITIDA")],
    ]
    xp.escribir_tabla(ws_ctl, "TControles", ["ID", "Control", "Estado"], ctl, 4, 2)
    fin_ctl = 4 + len(ctl)
    xp.definir_nombre(wb, "ModeloOK", f"'Controles'!$C${fin_ctl + 3}")
    ws_ctl[f"C{fin_ctl + 2}"] = "Estado global del modelo"
    ws_ctl[f"C{fin_ctl + 2}"].font = Font(bold=True)
    ws_ctl[f"C{fin_ctl + 3}"] = ('=IF(COUNTIF(TControles[Estado],"ERROR*")>0,"REVISAR: hay controles en error",'
                                 'IF(COUNTIF(TControles[Estado],"ADVERTENCIA*")>0,"OK CON ADVERTENCIAS","OK"))')
    ws_ctl[f"C{fin_ctl + 3}"].font = Font(bold=True, size=12)
    xp.semaforo(ws_ctl, f"D5:D{fin_ctl}")
    xp.semaforo(ws_ctl, f"C{fin_ctl + 3}:C{fin_ctl + 3}")
    ws_ctl.column_dimensions["B"].width = 8
    ws_ctl.column_dimensions["C"].width = 78
    ws_ctl.column_dimensions["D"].width = 52

    # ------------------------------------------------------------------ Resumen
    xp.cabecera_modelo(ws_res, "RiskLive · ¿Con qué probabilidad llego a mi meta?",
                       '="Aporte de $"&FIXED(AporteMensualCorrida,0)&" al mes durante "&HorizonteAnios&" años · portafolio "'
                       '&ROUND(PesoCDTCorrida*100,0)&"/"&ROUND(PesoCOLCAPCorrida*100,0)&"/"&ROUND(PesoUSDCorrida*100,0)'
                       '&" CDT/COLCAP/USD · "&NSim&" simulaciones"')
    ws_res["L2"] = "Estado del modelo"
    ws_res["L2"].font = Font(size=8, color="7F7F7F")
    ws_res["L3"] = "=ModeloOK"
    ws_res["L3"].font = Font(bold=True)
    xp.semaforo(ws_res, "L3:L3")
    xp.tarjeta_kpi(ws_res, "B6", "PROBABILIDAD DE LLEGAR A LA META", "=PMeta", xp.FMT_PCT)
    xp.tarjeta_kpi(ws_res, "D6", "VALOR REAL MEDIANO (COP)", "=MedianaReal", xp.FMT_COP)
    xp.tarjeta_kpi(ws_res, "F6", "PERCENTIL 5 (COP)", "=P5Real", xp.FMT_COP)
    xp.tarjeta_kpi(ws_res, "H6", "APORTE PARA 80 % (COP/MES)", "=Aporte80", xp.FMT_COP)
    xp.tarjeta_kpi(ws_res, "J6", "META (COP DE HOY)", "=MetaCorrida", xp.FMT_COP)
    xp.tarjeta_kpi(ws_res, "L6", "DATOS AL CORTE", '=YEAR(FechaDatos)&"-"&RIGHT("0"&MONTH(FechaDatos),2)&"-"&RIGHT("0"&DAY(FechaDatos),2)')
    ws_res["B9"] = "HALLAZGOS"
    ws_res["B9"].font = Font(bold=True, size=8, color=xp.GRIS_TENUE)
    ws_res["B10"] = ('="1. Con $"&FIXED(AporteMensualCorrida,0)&" al mes durante "&HorizonteAnios&" años, la probabilidad de '
                     'llegar a $"&FIXED(MetaCorrida,0)&" (pesos de hoy) es "&ROUND(PMeta*100,1)&" %."')
    ws_res["B11"] = ('="2. El 90 % de los resultados queda entre $"&FIXED(P5Real,0)&" y $"&FIXED(P95Real,0)&" en pesos de hoy; '
                     'el peor caso razonable (percentil 5) es $"&FIXED(P5Real,0)&"."')
    ws_res["B12"] = ('="3. Para tener 80 % de probabilidad hace falta aportar $"&FIXED(Aporte80,0)&" al mes, "&'
                     'IF(Aporte80>AporteMensualCorrida,"más","menos")&" que los $"&FIXED(AporteMensualCorrida,0)&" actuales."')
    n_pm = pm.shape[0]
    fan = LineChart()
    fan.title = None
    fan.y_axis.number_format = '#,##0,,"M"'
    fan.y_axis.delete = False
    fan.x_axis.delete = False
    fan.x_axis.tickLblSkip = max(1, n_pm // 10)
    for col_idx, color, ancho, dash in ((2, "99E0D6", 12000, None), (3, "4CC0B0", 15000, None), (4, "0F766E", 30000, None),
                                        (5, "4CC0B0", 15000, None), (6, "99E0D6", 12000, None), (7, ROJO, 19000, "dash")):
        fan.add_data(Reference(ws_pct, min_col=col_idx, min_row=4, max_row=4 + n_pm), titles_from_data=True)
        s = fan.series[-1]
        s.graphicalProperties.line.solidFill = color
        s.graphicalProperties.line.width = ancho
        s.smooth = False
        if dash:
            s.graphicalProperties.line.dashStyle = dash
    fan.set_categories(Reference(ws_pct, min_col=1, min_row=5, max_row=4 + n_pm))
    fan.legend.position = "t"
    _grafico_limpio(fan)
    fan.height, fan.width = 8.8, 16.5
    ws_res["B14"] = "SALDO REAL MES A MES (MILLONES DE COP DE HOY) · PERCENTILES 5, 25, 50, 75 Y 95 VS META"
    ws_res["B14"].font = Font(bold=True, size=8, color=xp.GRIS_TENUE)
    ws_res.add_chart(fan, "B15")
    hist = BarChart()
    hist.type = "col"
    hist.title = None
    hist.gapWidth = 15
    hist.y_axis.delete = False
    hist.x_axis.delete = False
    hist.legend = None
    hist.add_data(Reference(ws_rsl, min_col=5, min_row=rh, max_row=rh + len(cuentas)), titles_from_data=True)
    hist.set_categories(Reference(ws_rsl, min_col=2, min_row=rh + 1, max_row=rh + len(cuentas)))
    hist.x_axis.tickLblSkip = 3
    serie = hist.series[0]
    serie.graphicalProperties.solidFill = TEAL
    for i, f in enumerate(filas_h):
        pt = DataPoint(idx=i)
        pt.graphicalProperties.solidFill = TEAL if f[4] == "Sobre la meta" else AMBAR
        serie.dPt.append(pt)
    _grafico_limpio(hist)
    hist.height, hist.width = 8.8, 16.5
    ws_res["H14"] = "VALOR REAL FINAL (MILLONES DE COP DE HOY) · ÁMBAR: BAJO LA META · VERDE AZULADO: SOBRE LA META"
    ws_res["H14"].font = Font(bold=True, size=8, color=xp.GRIS_TENUE)
    ws_res.add_chart(hist, "H15")
    ws_res["B35"] = "SUPUESTOS DEL DÍA"
    ws_res["B35"].font = Font(bold=True, size=8, color=xp.GRIS_TENUE)
    for j, h in enumerate(["VARIABLE", "", "VALOR", "FUENTE"]):
        xp.encabezado(ws_res.cell(row=36, column=2 + [0, 1, 4, 6][j], value=h))
    for c_ in (3, 4, 6):
        xp.encabezado(ws_res.cell(row=36, column=c_))
    resumen_sup = ctx["supuestos"][:6]
    for i, s in enumerate(resumen_sup):
        r = 37 + i
        ws_res.cell(row=r, column=2, value=s["Variable"])
        c_ = ws_res.cell(row=r, column=6, value=f"=INDEX(TSupuestos[Valor],{i + 1})")
        c_.number_format = s["Formato"]
        c_.alignment = Alignment(horizontal="right")
        ws_res.cell(row=r, column=8, value=s["Fuente"])
        for k in range(2, 13):
            ws_res.cell(row=r, column=k).border = xp._BORDE
    xp.nota(ws_res, f"B{37 + len(resumen_sup) + 1}", "Educación financiera, no asesoría de inversión. Los resultados dependen de la "
            "calibración con datos históricos; ver hoja Fuentes y las limitaciones del informe.", 12, 32)
    ws_res.column_dimensions["A"].width = 3
    for col in "BCDEFGHIJKLM":
        ws_res.column_dimensions[col].width = 14

    for hoja, rol in ((ws_res, "salida"), (ws_ctl, "control"), (ws_par, "entrada"), (ws_sup, "dato"), (ws_cal, "dato"),
                      (ws_mod, "calculo"), (ws_sim, "calculo"), (ws_rsl, "salida"), (ws_pct, "dato"), (ws_fue, "dato")):
        xp.formato_hoja(hoja, rol)
    xp.estandarizar_fuentes(wb)
    wb.save(ruta)
