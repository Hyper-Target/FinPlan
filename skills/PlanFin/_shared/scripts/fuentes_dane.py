"""DANE: IPC total nacional, archivo de índices (serie de empalme, base dic-2018).

URL: https://www.dane.gov.co/files/operaciones/IPC/{mes}{AAAA}/anex-IPC-Indices-{mes}{AAAA}.xlsx
Hoja 'IndicesIPC': una fila de encabezado con 'Mes' en la columna A y los años a la derecha (2003, 2004…);
debajo, 12 filas (Enero … Diciembre). La última celda con número es el índice más reciente.
Inflación a 12 meses = índice(t) / índice(t-12) - 1.
"""
from __future__ import annotations

import io
from datetime import date

import openpyxl
import pandas as pd

from _http import FuenteError, get

MESES = ["ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic"]
NOMBRES_MES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto",
               "Septiembre", "Octubre", "Noviembre", "Diciembre"]


def url_indices(anio: int, mes: int) -> str:
    m = MESES[mes - 1]
    return f"https://www.dane.gov.co/files/operaciones/IPC/{m}{anio}/anex-IPC-Indices-{m}{anio}.xlsx"


def parsear_indices(contenido: bytes) -> pd.DataFrame:
    """Convierte el xlsx de índices en DataFrame [fecha (fin de mes), indice]."""
    wb = openpyxl.load_workbook(io.BytesIO(contenido), data_only=True)
    ws = wb["IndicesIPC"] if "IndicesIPC" in wb.sheetnames else wb[wb.sheetnames[0]]
    fila_enc = None
    for row in ws.iter_rows(min_row=1, max_row=30):
        if str(row[0].value).strip().lower() == "mes":
            fila_enc = row[0].row
            break
    if fila_enc is None:
        raise FuenteError("No se encontró la fila de encabezado 'Mes' en el archivo de IPC del DANE (formato cambió).")
    anios = {}
    for c in ws[fila_enc][1:]:
        try:
            anios[c.column] = int(str(c.value).strip()[:4])
        except (TypeError, ValueError):
            continue
    datos = []
    for i, nombre in enumerate(NOMBRES_MES):
        fila = ws[fila_enc + 1 + i]
        if str(fila[0].value).strip().lower() != nombre.lower():
            raise FuenteError(f"Se esperaba '{nombre}' bajo el encabezado del IPC y se encontró {fila[0].value!r}.")
        for c in fila[1:]:
            if c.column in anios and isinstance(c.value, (int, float)):
                fecha = pd.Timestamp(anios[c.column], i + 1, 1) + pd.offsets.MonthEnd(0)
                datos.append((fecha, float(c.value)))
    df = pd.DataFrame(datos, columns=["fecha", "indice"]).sort_values("fecha").reset_index(drop=True)
    if len(df) < 24:
        raise FuenteError("El archivo de IPC del DANE trae menos de 24 meses de datos; formato inesperado.")
    return df


def descargar_indices(hoy: date | None = None, max_meses_atras: int = 3) -> tuple[pd.DataFrame, str]:
    """Busca el archivo más reciente: prueba el mes actual y hasta `max_meses_atras` meses previos.

    Devuelve (DataFrame de índices, URL usada). Lanza FuenteError si ninguno existe.
    """
    hoy = hoy or date.today()
    anio, mes = hoy.year, hoy.month
    intentadas = []
    for _ in range(max_meses_atras + 1):
        url = url_indices(anio, mes)
        intentadas.append(url)
        r = get(url, nombre_fuente="DANE IPC")
        if r.status_code == 200 and r.content[:2] == b"PK":
            return parsear_indices(r.content), url
        mes -= 1
        if mes == 0:
            mes, anio = 12, anio - 1
    raise FuenteError("No se encontró el archivo de IPC del DANE. URLs probadas: " + " | ".join(intentadas))


def inflacion_12m(df: pd.DataFrame) -> dict:
    """Variación anual del último índice: {fecha, valor (fracción), indice, indice_previo}."""
    ult = df.iloc[-1]
    previo = df[df["fecha"] == ult["fecha"] - pd.offsets.MonthEnd(12)]
    if previo.empty:
        raise FuenteError("No hay índice de hace 12 meses para calcular la inflación anual.")
    ip = float(previo.iloc[0]["indice"])
    return {"fecha": ult["fecha"].date(), "valor": float(ult["indice"]) / ip - 1.0,
            "indice": float(ult["indice"]), "indice_previo": ip}


def inflacion_mensual(df: pd.DataFrame) -> pd.DataFrame:
    """Serie [fecha, inflacion_mensual] (fracción)."""
    out = df.copy()
    out["inflacion_mensual"] = out["indice"].pct_change()
    return out.dropna()[["fecha", "inflacion_mensual"]].reset_index(drop=True)
