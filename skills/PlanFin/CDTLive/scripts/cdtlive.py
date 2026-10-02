#!/usr/bin/env python
"""CDTLive: compara CDT de bancos y fintech de Colombia con datos oficiales del día.

Uso (ver SKILL.md):
  python cdtlive.py --monto 10000000 --plazo 360 --periodicidad vencimiento
  python cdtlive.py --rellenar RUTA.md --grupos "..." --inflacion "..." --ahorro "..." --alertas "..."

Entrega SIEMPRE un Excel con fórmulas vivas y un Markdown que lo interpreta.
"""
from __future__ import annotations

import argparse
import json
import re
import subprocess
import sys
import tempfile
from datetime import date, datetime
from pathlib import Path

import pandas as pd

SCRIPT_DIR = Path(__file__).resolve().parent
SKILL_DIR = SCRIPT_DIR.parent
PLANFIN = SKILL_DIR.parent
sys.path.insert(0, str(PLANFIN / "_shared" / "scripts"))

import excel_profe as xp  # noqa: E402
from informe import (f_cop, f_fecha, f_pct, f_pp, normalizar_numero, numeros_de, rellenar,  # noqa: E402,F401
                     tabla_md, validar_interpretacion)
import fuentes_banrep as br  # noqa: E402
import fuentes_dane as dn  # noqa: E402
import fuentes_sfc as sfc  # noqa: E402
import tasas as tx  # noqa: E402
from _http import ADVERTENCIAS, FuenteError  # noqa: E402

RAIZ = PLANFIN.parent.parent  # carpeta "Planeación Financiera"
SALIDA_DEFECTO = RAIZ / "salidas" / "CDTLive"

COBERTURA_FOGAFIN = 50_000_000  # COP por persona y por entidad (capital + intereses). Verificado 2026-09-29
FUENTE_FOGAFIN = "https://www.fogafin.gov.co/sites/default/files/2025-12/ABC%20del%20proceso%20de%20pago%20del%20seguro%20de%20dep%C3%B3sitos.pdf"
ADVERTENCIA = "Educación financiera, no asesoría de inversión. Verifique la tasa con la entidad antes de invertir."

PERIODICIDADES = {"vencimiento": None, "mensual": 30, "trimestral": 90, "semestral": 180, "anual": 360}
ETIQUETA_PERIODICIDAD = {"vencimiento": "Al vencimiento", "mensual": "Mensual", "trimestral": "Trimestral",
                         "semestral": "Semestral", "anual": "Anual"}
CATEGORIAS = ["Banco tradicional", "Banco digital", "Fintech"]
COLOR_CATEGORIA = {"Banco tradicional": "1F2937", "Banco digital": "0D9488", "Fintech": "F59E0B"}
MAX_PERIODOS_FLUJO = 60
ENTIDADES_AHORRO_MD = ["Nu Colombia", "Nequi", "RappiPay", "Lulo Bank", "Bancolombia", "Davivienda"]


# ----------------------------------------------------------------------------------------------------------------
# Cálculo (réplica en Python de las fórmulas del Excel; sirve para validar)
# ----------------------------------------------------------------------------------------------------------------
def dias_periodo_de(periodicidad: str, plazo: int) -> int:
    dp = PERIODICIDADES[periodicidad]
    return plazo if dp is None else dp


def calcular_cdt(tasa_ea: float, monto: float, plazo: int, dias_periodo: int, retencion: float = 0.04,
                 base: int = 365, ipc12m: float = 0.0, cobertura: float = COBERTURA_FOGAFIN) -> dict:
    """Cálculo de un CDT con intereses pagados (no capitalizados) cada `dias_periodo` días."""
    if plazo % dias_periodo != 0:
        raise ValueError(f"El plazo ({plazo} días) no es múltiplo del periodo de pago ({dias_periodo} días).")
    n = plazo // dias_periodo
    ip = tx.ea_a_periodica(tasa_ea, dias_periodo, base)
    bruto_periodo = monto * ip
    bruto_total = bruto_periodo * n
    ret = bruto_total * retencion
    neto = bruto_total - ret
    rent_neta_ea = (1 + ip * (1 - retencion)) ** (base / dias_periodo) - 1  # = TIR anualizada (bono a la par)
    rent_real_ea = tx.tasa_real(rent_neta_ea, ipc12m)
    valor_final = monto + neto
    return {
        "TasaEA": tasa_ea, "TasaPeriodica": ip, "NumPeriodos": n, "InteresBrutoPeriodo": bruto_periodo,
        "InteresBrutoTotal": bruto_total, "RetencionCOP": ret, "InteresNeto": neto, "ValorFinal": valor_final,
        "RentNetaEA": rent_neta_ea, "RentRealEA": rent_real_ea,
        "AlertaFogafin": "Excede cobertura Fogafín" if valor_final > cobertura else "",
    }


def flujo_neto(monto: float, plazo: int, dias_periodo: int, tasa_ea: float, retencion: float, base: int = 365) -> list[float]:
    """Flujo de caja neto del inversionista: -monto en 0, interés neto cada periodo, capital al final."""
    n = plazo // dias_periodo
    ip = tx.ea_a_periodica(tasa_ea, dias_periodo, base)
    flujos = [-monto]
    for t in range(1, n + 1):
        flujos.append(monto * ip * (1 - retencion) + (monto if t == n else 0.0))
    return flujos


def calcular_ahorro(tasa_ea: float, monto: float, plazo: int, retencion: float, base: int = 365) -> dict:
    bruto = monto * ((1 + tasa_ea) ** (plazo / base) - 1)
    return {"InteresBrutoAhorro": bruto, "InteresNetoAhorro": bruto * (1 - retencion)}


# ----------------------------------------------------------------------------------------------------------------
# Formato de números en español (1.234.567 y 11,50 %)
# ----------------------------------------------------------------------------------------------------------------
# ----------------------------------------------------------------------------------------------------------------
# Recolección de datos
# ----------------------------------------------------------------------------------------------------------------
def recolectar(monto: int, plazo: int, periodicidad: str, retencion: float, base: int) -> dict:
    ahora = datetime.now()
    entidades = pd.read_csv(SKILL_DIR / "entidades.csv", encoding="utf-8", dtype={"marca_comercial": str}).fillna("")
    fuentes: list[dict] = []

    df, cortes, url_sfc = sfc.descargar_captacion(6)
    fuentes.append({"Fuente": "Superfinanciera (datos.gov.co, dataset axk9-g2nh)", "URL": url_sfc,
                    "FechaHoraConsulta": ahora.strftime("%Y-%m-%d %H:%M"), "FilasRecibidas": len(df),
                    "Detalle": f"Cortes: {', '.join(c.isoformat() for c in cortes)}; ucas 1 (CDT) y 7 (ahorro)"})
    faltantes = sfc.entidades_faltantes(df, entidades)
    sub = sfc.plazo_a_subcuenta(plazo)
    tasas = sfc.seleccionar(df, entidades, 1, sub, cortes)
    ofrece = dict(zip(entidades["nombre_corto"], entidades["ofrece_cdt"]))
    tasas.loc[(tasas["Estado"] == "Sin dato") & (tasas["Entidad"].map(ofrece) == "no"), "Estado"] = "No emite CDT"
    ahorro = sfc.seleccionar(df, entidades, 7, 10, cortes)

    refs = []
    sistema = br.cdt_sistema_por_plazo(plazo)
    tpm, ibr, dtf, meta = (br.ultimo(br.SERIES[k]) for k in ("TPM", "IBR3M", "DTF90", "META_INFLACION"))
    cdt = {k: br.ultimo(br.SERIES[k]) for k in ("CDT90", "CDT180", "CDT360")}
    for u in [cdt["CDT90"], cdt["CDT180"], cdt["CDT360"], dtf, ibr, tpm, meta]:
        refs.append({"Serie": u["nombre"], "IDSerie": u["id"], "Fecha": u["fecha"], "ValorPct": u["valor"] / 100,
                     "Unidad": u["unidad"], "Fuente": u["url"]})
        fuentes.append({"Fuente": f"BanRep Suameca serie {u['id']}", "URL": u["url"],
                        "FechaHoraConsulta": ahora.strftime("%Y-%m-%d %H:%M"), "FilasRecibidas": None,
                        "Detalle": u["nombre"]})

    ipc_nota = ""
    try:
        idx, url_ipc = dn.descargar_indices()
        ipc = dn.inflacion_12m(idx)
        ipc12 = ipc["valor"]
        refs.append({"Serie": "IPC total nacional, variación anual (DANE)", "IDSerie": "", "Fecha": ipc["fecha"],
                     "ValorPct": ipc12, "Unidad": "% anual", "Fuente": url_ipc})
        fuentes.append({"Fuente": "DANE IPC (archivo de índices)", "URL": url_ipc,
                        "FechaHoraConsulta": ahora.strftime("%Y-%m-%d %H:%M"), "FilasRecibidas": len(idx),
                        "Detalle": f"Índice {ipc['indice']} / {ipc['indice_previo']} - 1"})
        fecha_ipc = ipc["fecha"]
    except FuenteError as e:
        ipc12 = meta["valor"] / 100
        fecha_ipc = meta["fecha"]
        ipc_nota = f"IPC no disponible, se usa la meta de inflación del BanRep ({f_pct(ipc12)}). Detalle: {e}"
        ADVERTENCIAS.append(ipc_nota)

    return {
        "ahora": ahora, "entidades": entidades, "cortes": cortes, "faltantes": faltantes, "subcuenta": sub,
        "tasas": tasas, "ahorro": ahorro, "sistema": sistema, "tpm": tpm, "ibr": ibr, "dtf": dtf, "meta": meta,
        "refs": refs, "ipc12m": ipc12, "fecha_ipc": fecha_ipc, "ipc_nota": ipc_nota, "fuentes": fuentes,
        "monto": monto, "plazo": plazo, "periodicidad": periodicidad, "retencion": retencion, "base": base,
    }


def calcular_todo(d: dict) -> dict:
    """Aplica el cálculo a cada entidad con dato y arma tablas para Excel y MD."""
    dp = dias_periodo_de(d["periodicidad"], d["plazo"])
    d["dias_periodo"] = dp
    ofrece = dict(zip(d["entidades"]["nombre_corto"], d["entidades"]["ofrece_cdt"]))
    ok = d["tasas"][(d["tasas"]["Estado"] == "OK") & (d["tasas"]["Entidad"].map(ofrece) == "si")]
    prom_sis = d["sistema"]["valor"] / 100
    filas = []
    for r in ok.itertuples():
        c = calcular_cdt(r.TasaEA, d["monto"], d["plazo"], dp, d["retencion"], d["base"], d["ipc12m"])
        c.update(Entidad=r.Entidad, Categoria=r.Categoria, MarcaComercial=r.MarcaComercial,
                 DifVsSistema=r.TasaEA - prom_sis, RezagoCortes=r.RezagoCortes, Alerta=r.Alerta)
        filas.append(c)
    sim = pd.DataFrame(filas).sort_values("RentNetaEA", ascending=False).reset_index(drop=True)
    sim["Ranking"] = sim["RentNetaEA"].rank(ascending=False, method="min").astype(int)
    d["sim"] = sim
    ah = d["ahorro"][d["ahorro"]["Estado"] == "OK"].copy()
    ah = ah.sort_values("TasaEA", ascending=False).reset_index(drop=True)
    for k in ("InteresBrutoAhorro", "InteresNetoAhorro"):
        ah[k] = 0.0
    for i, r in ah.iterrows():
        a = calcular_ahorro(r["TasaEA"], d["monto"], d["plazo"], d["retencion"], d["base"])
        ah.loc[i, "InteresBrutoAhorro"] = a["InteresBrutoAhorro"]
        ah.loc[i, "InteresNetoAhorro"] = a["InteresNetoAhorro"]
    d["ah"] = ah
    if sim.empty:
        raise FuenteError("Ninguna entidad tiene dato de CDT para ese plazo en los últimos 6 cortes.")
    mejor = sim.iloc[0]
    d["mejor"] = mejor
    ah["DifVsMejorCDT"] = ah["InteresNetoAhorro"] - mejor["InteresNeto"]
    # Validación de flujo: ΣXfd = 0 a la TIR y TIR anualizada = RentNetaEA de la mejor
    fl = flujo_neto(d["monto"], d["plazo"], dp, mejor["TasaEA"], d["retencion"], d["base"])
    tir_p = tx.tir(fl)
    d["flujo"] = fl
    d["tir_per"] = tir_p
    d["tir_ea"] = tx.periodica_a_ea(tir_p, dp, d["base"])
    d["suma_xfd"] = sum(f / (1 + tir_p) ** t for t, f in enumerate(fl))
    return d


def validaciones_python(d: dict) -> list[dict]:
    hoy = date.today()
    antig = (hoy - d["cortes"][0]).days
    n = len(d["sim"])
    plazo_ok = d["plazo"] % d["dias_periodo"] == 0 and d["plazo"] // d["dias_periodo"] <= MAX_PERIODOS_FLUJO
    return [
        {"Validacion": "ΣXfd = 0 a la TIR (tolerancia $1)", "Resultado": "OK" if abs(d["suma_xfd"]) < 1 else "ERROR",
         "Detalle": f"ΣXfd = {d['suma_xfd']:.6f}"},
        {"Validacion": "TIR anualizada = rentabilidad neta EA de la mejor opción",
         "Resultado": "OK" if abs(d["tir_ea"] - d["mejor"]["RentNetaEA"]) < 1e-9 else "ERROR",
         "Detalle": f"TIR EA {d['tir_ea']:.8f} vs {d['mejor']['RentNetaEA']:.8f}"},
        {"Validacion": "Fecha de corte con menos de 7 días", "Resultado": "OK" if antig < 7 else "ADVERTENCIA",
         "Detalle": f"Último corte {d['cortes'][0]} ({antig} días)"},
        {"Validacion": "Al menos 10 entidades con dato", "Resultado": "OK" if n >= 10 else "ADVERTENCIA",
         "Detalle": f"{n} entidades con dato de CDT"},
        {"Validacion": "Plazo múltiplo del periodo y hasta 60 periodos", "Resultado": "OK" if plazo_ok else "ERROR",
         "Detalle": f"{d['plazo']} días / {d['dias_periodo']} días por periodo"},
        {"Validacion": "Entidades de entidades.csv presentes en el dataset",
         "Resultado": "OK" if not d["faltantes"] else "ADVERTENCIA",
         "Detalle": "Todas presentes" if not d["faltantes"] else "Faltan: " + ", ".join(d["faltantes"])},
    ]


# ----------------------------------------------------------------------------------------------------------------
# Excel
# ----------------------------------------------------------------------------------------------------------------
def construir_excel(d: dict, ruta: Path, checks: list[dict]) -> None:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.chart.series import DataPoint
    from openpyxl.chart.shapes import GraphicalProperties
    from openpyxl.drawing.line import LineProperties
    from openpyxl.styles import Font
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws_res = wb.active
    ws_res.title = "Resumen"
    ws_ctl = wb.create_sheet("Controles")
    ws_par = wb.create_sheet("Parametros")
    ws_tas = wb.create_sheet("Tasas")
    ws_sim = wb.create_sheet("Simulacion")
    ws_flu = wb.create_sheet("Flujo_Mejor")
    ws_aho = wb.create_sheet("Ahorro")
    ws_ref = wb.create_sheet("Referencias")
    ws_fue = wb.create_sheet("Fuentes")
    wb.calculation.fullCalcOnLoad = True

    # ---- Parametros (nombres definidos, celdas azules = entradas) ----
    xp.titulo(ws_par, "B2", "Parámetros del CDT (celdas azules: cámbielas y todo se recalcula)")
    for j, h in enumerate(["Variable", "Valor", "Unidad", "Observación"]):
        xp.encabezado(ws_par.cell(row=4, column=2 + j, value=h))
    P = xp.fila_parametro
    P(ws_par, 5, "Monto", d["monto"], "COP", "Capital a invertir", "Monto", wb, xp.FMT_COP, True)
    P(ws_par, 6, "Plazo", d["plazo"], "días", "Plazo del CDT", "PlazoDias", wb, "0", True)
    P(ws_par, 7, "Periodicidad de pago", ETIQUETA_PERIODICIDAD[d["periodicidad"]], "", "Cómo se pagan los intereses",
      "Periodicidad", wb, None, True)
    dv = DataValidation(type="list", formula1='"Al vencimiento,Mensual,Trimestral,Semestral,Anual"', allow_blank=False)
    ws_par.add_data_validation(dv)
    dv.add("C7")
    P(ws_par, 8, "Días por periodo", '=IF(Periodicidad="Mensual",30,IF(Periodicidad="Trimestral",90,'
      'IF(Periodicidad="Semestral",180,IF(Periodicidad="Anual",360,PlazoDias))))', "días", "Calculado", "DiasPeriodo", wb, "0")
    P(ws_par, 9, "Número de periodos", "=PlazoDias/DiasPeriodo", "", "Calculado", "NPer", wb, "0.00")
    P(ws_par, 10, "Chequeo del plazo", f'=IF(AND(MOD(PlazoDias,DiasPeriodo)=0,NPer<={MAX_PERIODOS_FLUJO}),"OK",'
      '"ERROR: plazo no múltiplo del periodo")', "", f"Hasta {MAX_PERIODOS_FLUJO} periodos en la hoja Flujo_Mejor",
      "ChequeoPlazo", wb)
    P(ws_par, 11, "Base de días del año", d["base"], "días", "Convención de la EA", "BaseDias", wb, "0", True)
    P(ws_par, 12, "Retención en la fuente", d["retencion"], "% sobre intereses", "Verifique la tarifa vigente con la entidad",
      "Retencion", wb, xp.FMT_PCT, True)
    P(ws_par, 13, "Inflación a 12 meses (IPC)", d["ipc12m"], "% anual", f"DANE, corte {f_fecha(d['fecha_ipc'])}" +
      (" (ver Fuentes: se usó la meta)" if d["ipc_nota"] else ""), "IPC12m", wb, xp.FMT_PCT, True)
    P(ws_par, 14, "Promedio del sistema (CDT)", d["sistema"]["valor"] / 100, "% EA bruta",
      f"BanRep serie {d['sistema']['id']}, {f_fecha(d['sistema']['fecha'])}", "PromedioSistema", wb, xp.FMT_PCT, True)
    P(ws_par, 15, "Cobertura seguro de depósitos", COBERTURA_FOGAFIN, "COP por persona y entidad",
      "Fogafín (capital + intereses). Verificado 2026-09-29", "CoberturaFogafin", wb, xp.FMT_COP, True)
    P(ws_par, 16, "Fecha del último corte", d["cortes"][0], "", "Superfinanciera, con un día de rezago", "FechaCorte", wb,
      xp.FMT_FECHA)
    P(ws_par, 17, "Fecha de consulta", d["ahora"].strftime("%Y-%m-%d %H:%M"), "", "", "FechaConsulta", wb)
    P(ws_par, 18, "Tasa de política monetaria", d["tpm"]["valor"] / 100, "% EA", f"BanRep, {f_fecha(d['tpm']['fecha'])}",
      "TPM", wb, xp.FMT_PCT)
    P(ws_par, 19, "IBR a 3 meses", d["ibr"]["valor"] / 100, "% nominal", f"BanRep, {f_fecha(d['ibr']['fecha'])}", "IBR3M",
      wb, xp.FMT_PCT)
    xp.ajustar_anchos(ws_par)
    ws_par.column_dimensions["E"].width = 55

    # ---- Tasas (datos crudos de la Superfinanciera, todas las entidades) ----
    xp.titulo(ws_tas, "A2", f"Tasas de captación por entidad, {sfc.DESC_SUBCUENTA.get(d['subcuenta'], '')} "
              f"(subcuenta {d['subcuenta']}): tasa promedio ponderada pagada, no tasa de cartelera")
    cols_t = ["Categoria", "Entidad", "MarcaComercial", "TipoEntidad", "CodigoEntidad", "Subcuenta", "PlazoDescripcion",
              "TasaEA", "MontoCaptadoMiles", "FechaCorte", "RezagoCortes", "Estado", "Alerta", "Fuente", "IDDataset",
              "ExtraidoEn"]
    filas_t = []
    for _, r in d["tasas"].iterrows():
        fila = [None if pd.isna(r[c]) else r[c] for c in cols_t[:13]]
        filas_t.append(fila + ["Superfinanciera de Colombia", "datos.gov.co/axk9-g2nh (uca 1)",
                               d["ahora"].strftime("%Y-%m-%d %H:%M")])
    xp.escribir_tabla(ws_tas, "TTasas", cols_t, filas_t, 4, 1,
                      {"TasaEA": xp.FMT_PCT, "MontoCaptadoMiles": xp.FMT_COP, "FechaCorte": xp.FMT_FECHA})
    ws_tas.freeze_panes = "A5"
    xp.ajustar_anchos(ws_tas)

    # ---- Simulacion (fórmulas estructuradas) ----
    sim = d["sim"]
    xp.titulo(ws_sim, "A2", "Simulación por entidad (todas las columnas de cálculo son fórmulas vivas)")
    cols_s = ["Entidad", "Categoria", "MarcaComercial", "TasaEA", "TasaPeriodica", "NumPeriodos", "InteresBrutoPeriodo",
              "InteresBrutoTotal", "RetencionCOP", "InteresNeto", "ValorFinal", "RentNetaEA", "RentRealEA",
              "DifVsSistema", "LineaSistemaNeta", "AlertaFogafin", "Ranking"]
    TR = "TSimulacion[[#This Row],[{}]]".format
    filas_s = []
    for r in sim.itertuples():
        filas_s.append([
            r.Entidad, r.Categoria, r.MarcaComercial,
            f"=INDEX(TTasas[TasaEA],MATCH({TR('Entidad')},TTasas[Entidad],0))",
            f"=(1+{TR('TasaEA')})^(DiasPeriodo/BaseDias)-1",
            "=NPer",
            f"=Monto*{TR('TasaPeriodica')}",
            f"={TR('InteresBrutoPeriodo')}*{TR('NumPeriodos')}",
            f"={TR('InteresBrutoTotal')}*Retencion",
            f"={TR('InteresBrutoTotal')}-{TR('RetencionCOP')}",
            f"=Monto+{TR('InteresNeto')}",
            f"=(1+{TR('TasaPeriodica')}*(1-Retencion))^(BaseDias/DiasPeriodo)-1",
            f"=(1+{TR('RentNetaEA')})/(1+IPC12m)-1",
            f"={TR('TasaEA')}-PromedioSistema",
            "=(1+((1+PromedioSistema)^(DiasPeriodo/BaseDias)-1)*(1-Retencion))^(BaseDias/DiasPeriodo)-1",
            f'=IF({TR("ValorFinal")}>CoberturaFogafin,"Excede cobertura Fogafín","")',
            f"=RANK({TR('RentNetaEA')},TSimulacion[RentNetaEA],0)",
        ])
    fmt_s = {"TasaEA": xp.FMT_PCT, "TasaPeriodica": "0.0000%", "NumPeriodos": "0", "InteresBrutoPeriodo": xp.FMT_COP,
             "InteresBrutoTotal": xp.FMT_COP, "RetencionCOP": xp.FMT_COP, "InteresNeto": xp.FMT_COP,
             "ValorFinal": xp.FMT_COP, "RentNetaEA": xp.FMT_PCT, "RentRealEA": xp.FMT_PCT, "DifVsSistema": "+0.00%;-0.00%",
             "LineaSistemaNeta": xp.FMT_PCT, "Ranking": "0"}
    xp.escribir_tabla(ws_sim, "TSimulacion", cols_s, filas_s, 4, 1, fmt_s)
    ws_sim.freeze_panes = "B5"
    xp.ajustar_anchos(ws_sim, 12, 26)
    n = len(sim)
    xp.nota(ws_sim, f"A{5 + n + 2}",
            "RentNetaEA = TIR anualizada del flujo neto (interés menos retención; supone reinversión de los intereses "
            "a la misma tasa). DifVsSistema = tasa bruta de la entidad menos el promedio bruto del sistema (BanRep). "
            "LineaSistemaNeta = el promedio del sistema pasado por el mismo cálculo neto, para comparar peras con peras.",
            10, 48)

    # ---- Flujo_Mejor (convención del profe: Period, Inflow, Outflow, FCN, fd, Xfd; ΣXfd = 0) ----
    xp.titulo(ws_flu, "B2", "Flujo de caja de la mejor opción y validación ΣXfd = 0")
    q = "MATCH(1,TSimulacion[Ranking],0)"
    items = [
        (4, "Mejor entidad (mayor rentabilidad neta)", f"=INDEX(TSimulacion[Entidad],{q})", "MejorEntidad", None),
        (5, "Tasa EA bruta de la mejor", f"=INDEX(TSimulacion[TasaEA],{q})", "TasaMejor", xp.FMT_PCT),
        (6, "Tasa periódica de la mejor", "=(1+TasaMejor)^(DiasPeriodo/BaseDias)-1", "TasaPerMejor", "0.0000%"),
        (7, "TIR periódica del flujo neto (IRR de FCN)", "=IRR(TFlujo[FCN],TasaPerMejor*(1-Retencion))", "TIRPer", "0.0000%"),
        (8, "TIR anualizada (EA)", "=(1+TIRPer)^(BaseDias/DiasPeriodo)-1", "TIREA", xp.FMT_PCT),
        (9, "Rentabilidad neta EA (hoja Simulacion)", f"=INDEX(TSimulacion[RentNetaEA],{q})", "RentNetaMejor", xp.FMT_PCT),
        (10, "ΣXfd", "=SUM(TFlujo[Xfd])", "SumXfd", "0.000000"),
        (11, "Validación ΣXfd = 0", '=IF(ABS(SumXfd)<1,"OK","ERROR")', "ValidaXfd", None),
        (12, "TIR coincide con la hoja Simulacion", '=IF(ABS(TIREA-RentNetaMejor)<0.000001,"OK","ERROR")', "ValidaTIR", None),
        (13, "Interés neto de la mejor opción", f"=INDEX(TSimulacion[InteresNeto],{q})", "MejorInteresNeto", xp.FMT_COP),
        (14, "Rentabilidad real EA de la mejor", f"=INDEX(TSimulacion[RentRealEA],{q})", "RentRealMejor", xp.FMT_PCT),
        (15, "Diferencia de tasa vs sistema (mejor)", f"=INDEX(TSimulacion[DifVsSistema],{q})", "DifSistemaMejor", "+0.00%;-0.00%"),
    ]
    for fila, etiqueta, formula, nombre, fmt in items:
        xp.fila_parametro(ws_flu, fila, etiqueta, formula, "", "", nombre, wb, fmt)
    TF = "TFlujo[[#This Row],[{}]]".format
    filas_f = []
    for per in range(0, MAX_PERIODOS_FLUJO + 1):
        filas_f.append([
            per,
            f"=IF({TF('Period')}=0,0,IF({TF('Period')}<=NPer,Monto*TasaPerMejor*(1-Retencion),0)"
            f"+IF({TF('Period')}=NPer,Monto,0))",
            f"=IF({TF('Period')}=0,Monto,0)",
            f"={TF('Inflow')}-{TF('Outflow')}",
            f"=(1+TIRPer)^(-{TF('Period')})",
            f"={TF('FCN')}*{TF('fd')}",
        ])
    xp.escribir_tabla(ws_flu, "TFlujo", ["Period", "Inflow", "Outflow", "FCN", "fd", "Xfd"], filas_f, 18, 2,
                      {"Inflow": xp.FMT_NUM, "Outflow": xp.FMT_NUM, "FCN": xp.FMT_NUM, "fd": "0.000000", "Xfd": "0.000000"})
    xp.ajustar_anchos(ws_flu, 12, 45)
    ws_flu.column_dimensions["B"].width = 42
    ws_flu.column_dimensions["C"].width = 22

    # ---- Ahorro (uca 7, subcuenta 10) ----
    ah = d["ah"]
    xp.titulo(ws_aho, "A2", "Cuentas de ahorro (persona natural) frente al mejor CDT, mismo monto y plazo")
    cols_a = ["Entidad", "Categoria", "TasaEA", "SaldoMiles", "FechaCorte", "InteresBrutoAhorro", "InteresNetoAhorro",
              "DifVsMejorCDT"]
    TA = "TAhorro[[#This Row],[{}]]".format
    filas_a = []
    for r in ah.itertuples():
        filas_a.append([
            r.Entidad, r.Categoria, r.TasaEA, r.MontoCaptadoMiles, r.FechaCorte,
            f"=Monto*((1+{TA('TasaEA')})^(PlazoDias/BaseDias)-1)",
            f"={TA('InteresBrutoAhorro')}*(1-Retencion)",
            f"={TA('InteresNetoAhorro')}-MejorInteresNeto",
        ])
    xp.escribir_tabla(ws_aho, "TAhorro", cols_a, filas_a, 4, 1,
                      {"TasaEA": xp.FMT_PCT, "SaldoMiles": xp.FMT_COP, "FechaCorte": xp.FMT_FECHA,
                       "InteresBrutoAhorro": xp.FMT_COP, "InteresNetoAhorro": xp.FMT_COP, "DifVsMejorCDT": '#,##0;[Red]-#,##0'})
    xp.nota(ws_aho, f"A{5 + len(ah) + 2}",
            "Ahorro, no CDT: la tasa es el promedio ponderado del saldo de depósitos de ahorro activos de persona natural "
            "(uca 7, subcuenta 10). Se usa la misma retención del CDT para comparar; el tratamiento tributario real del "
            "ahorro puede diferir. DifVsMejorCDT negativo = lo que se deja de ganar frente al mejor CDT.", 8, 60)
    xp.ajustar_anchos(ws_aho, 12, 30)

    # ---- Referencias ----
    xp.titulo(ws_ref, "A2", "Series de referencia (BanRep y DANE)")
    cols_r = ["Serie", "IDSerie", "Fecha", "Valor", "Unidad", "Fuente"]
    filas_r = [[r["Serie"], r["IDSerie"], r["Fecha"], r["ValorPct"], r["Unidad"], r["Fuente"]] for r in d["refs"]]
    xp.escribir_tabla(ws_ref, "TReferencias", cols_r, filas_r, 4, 1, {"Fecha": xp.FMT_FECHA, "Valor": xp.FMT_PCT})
    xp.ajustar_anchos(ws_ref, 12, 70)

    # ---- Fuentes ----
    xp.titulo(ws_fue, "A2", "Fuentes, validaciones y notas metodológicas")
    cols_f = ["Fuente", "URL", "FechaHoraConsulta", "FilasRecibidas", "Detalle"]
    filas_fu = [[f["Fuente"], f["URL"], f["FechaHoraConsulta"], f["FilasRecibidas"], f["Detalle"]] for f in d["fuentes"]]
    xp.escribir_tabla(ws_fue, "TFuentes", cols_f, filas_fu, 4, 1)
    r0 = 4 + len(filas_fu) + 3
    xp.escribir_tabla(ws_fue, "TValidaciones", ["Validacion", "Resultado", "Detalle"],
                      [[c["Validacion"], c["Resultado"], c["Detalle"]] for c in checks], r0, 1)
    r1 = r0 + len(checks) + 3
    sin_dato = d["tasas"][d["tasas"]["Estado"] != "OK"]
    notas = [
        "Qué es la tasa: promedio ponderado por monto de lo que cada entidad pagó por CDT emitidos en el corte, en % EA. "
        "No es la tasa de cartelera ni la que le ofrezcan en la oficina.",
        "Por qué no se usa el simulador de Davivienda: el sitio tiene protección antibots y su backend respondió con "
        "error 500. Se replicó su cálculo en Python aplicado a la tasa que Davivienda reporta a la Superfinanciera.",
        f"Cobertura del seguro de depósitos: ${COBERTURA_FOGAFIN:,.0f} por persona y por entidad (capital más intereses). "
        f"Fuente: {FUENTE_FOGAFIN}",
        "Entidades sin dato o que no emiten CDT en este plazo: " + (", ".join(f"{r.Entidad} ({r.Estado})" for r in sin_dato.itertuples()) or "ninguna"),
        "Entidades con rezago: " + (", ".join(f"{r.Entidad} ({int(r.RezagoCortes)} corte)" for r in d["tasas"].itertuples()
                                              if r.Estado == "OK" and r.RezagoCortes and r.RezagoCortes > 0) or "ninguna"),
    ] + [f"ADVERTENCIA: {a}" for a in ADVERTENCIAS] + [ADVERTENCIA]
    for i, t in enumerate(notas):
        xp.nota(ws_fue, f"A{r1 + i}", t, 5, 34)
    xp.ajustar_anchos(ws_fue, 14, 60)
    ws_fue.column_dimensions["B"].width = 70

    # ---- Controles (integridad del modelo, todo con fórmulas vivas) ----
    xp.titulo(ws_ctl, "B2", "Controles de integridad del modelo", 16)
    ctl = [
        ["C01", "Suma de flujos descontados a la TIR es cero (ΣXfd)", '=IF(ABS(SumXfd)<1,"OK","ERROR")'],
        ["C02", "TIR anualizada del flujo coincide con la rentabilidad neta de la hoja Simulacion", '=IF(ValidaTIR="OK","OK","ERROR")'],
        ["C03", "El plazo es múltiplo del periodo de pago y cabe en la hoja de flujos", '=IF(ChequeoPlazo="OK","OK","ERROR")'],
        ["C04", "Las celdas de cálculo no contienen errores de fórmula",
         '=IF(SUMPRODUCT(--ISERROR(TSimulacion[[TasaEA]:[Ranking]]))+SUMPRODUCT(--ISERROR(TAhorro[[InteresBrutoAhorro]:[DifVsMejorCDT]]))=0,"OK","ERROR")'],
        ["C05", "Todas las tasas de la simulación están entre 0 % y 25 % EA",
         '=IF(AND(MIN(TSimulacion[TasaEA])>0,MAX(TSimulacion[TasaEA])<=0.25),"OK","ERROR")'],
        ["C06", "Hay un único primer lugar en el ranking",
         '=IF(COUNTIF(TSimulacion[Ranking],1)=1,"OK","ADVERTENCIA: empate en el primer lugar")'],
        ["C07", "Al menos 10 entidades con dato",
         '=IF(COUNTA(TSimulacion[Entidad])>=10,"OK","ADVERTENCIA: menos de 10 entidades")'],
        ["C08", "Dato de la Superfinanciera con menos de 7 días de antigüedad",
         '=IF(TODAY()-FechaCorte<7,"OK","ADVERTENCIA: dato con "&(TODAY()-FechaCorte)&" días")'],
        ["C09", "El valor final está dentro de la cobertura del seguro de depósitos",
         '=IF(INDEX(TSimulacion[ValorFinal],1)<=CoberturaFogafin,"OK","ADVERTENCIA: excede la cobertura de Fogafín")'],
        ["C10", "Conciliación del Excel contra un cálculo independiente en Python", d.get("estado_conciliacion", "OMITIDA")],
    ]
    xp.escribir_tabla(ws_ctl, "TControles", ["ID", "Control", "Estado"], ctl, 4, 2)
    xp.definir_nombre(wb, "ModeloOK", "'Controles'!$C$17")
    ws_ctl["C16"] = "Estado global del modelo"
    ws_ctl["C16"].font = Font(bold=True)
    ws_ctl["C17"] = ('=IF(COUNTIF(TControles[Estado],"ERROR*")>0,"REVISAR: hay controles en error",'
                     'IF(COUNTIF(TControles[Estado],"ADVERTENCIA*")>0,"OK CON ADVERTENCIAS","OK"))')
    ws_ctl["C17"].font = Font(bold=True, size=12)
    ws_ctl["C17"].alignment = xp.Alignment(horizontal="left")
    xp.semaforo(ws_ctl, "D5:D14")
    xp.semaforo(ws_ctl, "C17:C17")
    ws_ctl.column_dimensions["B"].width = 8
    ws_ctl.column_dimensions["C"].width = 78
    ws_ctl.column_dimensions["D"].width = 30

    # ---- Resumen ejecutivo ----
    xp.cabecera_modelo(ws_res, "CDTLive · ¿Dónde rinde más un CDT hoy en Colombia?",
                       '="CDT de $"&FIXED(Monto,0)&" COP a "&PlazoDias&" días  ·  intereses: "&Periodicidad&"  ·  datos al corte "&YEAR(FechaCorte)&"-"&RIGHT("0"&MONTH(FechaCorte),2)&"-"&RIGHT("0"&DAY(FechaCorte),2)')
    ws_res["K2"] = "Estado del modelo"
    ws_res["K2"].font = Font(size=8, color="7F7F7F")
    ws_res["K3"] = "=ModeloOK"
    ws_res["K3"].font = Font(bold=True)
    xp.semaforo(ws_res, "K3:K3")
    xp.tarjeta_kpi(ws_res, "B6", "MEJOR ENTIDAD", "=MejorEntidad", None, 2)
    xp.tarjeta_kpi(ws_res, "D6", "TASA EA BRUTA", "=TasaMejor", xp.FMT_PCT)
    xp.tarjeta_kpi(ws_res, "F6", "INTERÉS NETO (COP)", "=MejorInteresNeto", xp.FMT_COP)
    xp.tarjeta_kpi(ws_res, "H6", "RENTABILIDAD REAL EA", "=RentRealMejor", xp.FMT_PCT)
    xp.tarjeta_kpi(ws_res, "J6", "VS PROMEDIO DEL SISTEMA", "=DifSistemaMejor", "[Color10]+0.00%;[Red]-0.00%")
    xp.tarjeta_kpi(ws_res, "L6", "CONTROL ΣXfd = 0", "=ValidaXfd")
    ws_res["B9"] = "HALLAZGOS"
    ws_res["B9"].font = Font(bold=True, size=8, color=xp.GRIS_TENUE)
    ws_res["B10"] = ('="1. "&MejorEntidad&" ofrece la mayor rentabilidad neta: "&ROUND(TasaMejor*100,2)&" % EA bruta, "'
                     '&ROUND(RentNetaMejor*100,2)&" % EA tras retención."')
    ws_res["B11"] = ('="2. Frente al promedio del sistema ("&ROUND(PromedioSistema*100,2)&" % EA), la mejor tasa está "'
                     '&ROUND(ABS(DifSistemaMejor)*100,2)&" puntos porcentuales "&IF(DifSistemaMejor>=0,"por encima.","por debajo.")')
    ws_res["B12"] = ('="3. Con inflación de "&ROUND(IPC12m*100,2)&" %, la rentabilidad real de la mejor opción es "'
                     '&ROUND(RentRealMejor*100,2)&" % EA."')
    ws_res["B14"] = "RANKING · 10 MEJORES POR RENTABILIDAD NETA"
    ws_res["B14"].font = Font(bold=True, size=8, color=xp.GRIS_TENUE)
    enc = ["#", "ENTIDAD", "CATEGORÍA", "TASA EA BRUTA", "RENT. NETA EA", "INTERÉS NETO (COP)", "RENT. REAL EA", "ALERTA"]
    cols_ref = [None, "Entidad", "Categoria", "TasaEA", "RentNetaEA", "InteresNeto", "RentRealEA", "AlertaFogafin"]
    fmts_ref = [None, None, None, xp.FMT_PCT, xp.FMT_PCT, xp.FMT_COP, xp.FMT_PCT, None]
    for j, h in enumerate(enc):
        cel = ws_res.cell(row=15, column=2 + j, value=h)
        xp.encabezado(cel)
        if fmts_ref[j]:
            cel.alignment = xp.Alignment(horizontal="right", vertical="center", wrap_text=True)
    for k in range(1, min(10, n) + 1):
        rr = 15 + k
        ws_res.cell(row=rr, column=2, value=k).border = xp._BORDE
        for j in range(1, len(enc)):
            c = ws_res.cell(row=rr, column=2 + j, value=f"=INDEX(TSimulacion[{cols_ref[j]}],{k})")
            c.border = xp._BORDE
            if fmts_ref[j]:
                c.number_format = fmts_ref[j]
    xp.barras_datos(ws_res, f"F16:F{15 + min(10, n)}")
    ws_res["B28"] = "POR TIPO DE ENTIDAD"
    ws_res["B28"].font = Font(bold=True, size=8, color=xp.GRIS_TENUE)
    for j, h in enumerate(["CATEGORÍA", "ENTIDADES", "TASA EA PROM.", "RENT. NETA EA PROM."]):
        xp.encabezado(ws_res.cell(row=29, column=3 + j, value=h))
    for i, cat in enumerate(CATEGORIAS):
        rr = 30 + i
        ws_res.cell(row=rr, column=3, value=cat).border = xp._BORDE
        ws_res.cell(row=rr, column=4, value=f"=COUNTIF(TSimulacion[Categoria],C{rr})").border = xp._BORDE
        c = ws_res.cell(row=rr, column=5, value=f'=IFERROR(AVERAGEIFS(TSimulacion[TasaEA],TSimulacion[Categoria],C{rr}),"Sin dato")')
        c.number_format = xp.FMT_PCT
        c.border = xp._BORDE
        c = ws_res.cell(row=rr, column=6, value=f'=IFERROR(AVERAGEIFS(TSimulacion[RentNetaEA],TSimulacion[Categoria],C{rr}),"Sin dato")')
        c.number_format = xp.FMT_PCT
        c.border = xp._BORDE
    ws_res.column_dimensions["A"].width = 3
    ws_res.column_dimensions["B"].width = 6
    ws_res.column_dimensions["C"].width = 26
    ws_res.column_dimensions["D"].width = 18
    for col in "EFGHIJKLM":
        ws_res.column_dimensions[col].width = 16

    top = min(15, n)
    bar = BarChart()
    bar.type = "col"
    bar.title = None
    bar.y_axis.number_format = "0.0%"
    bar.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="E5E7EB", w=6350))
    bar.x_axis.delete = False
    bar.y_axis.delete = False
    bar.add_data(Reference(ws_sim, min_col=12, min_row=4, max_row=4 + top), titles_from_data=True)
    bar.set_categories(Reference(ws_sim, min_col=1, min_row=5, max_row=4 + top))
    s = bar.series[0]
    s.graphicalProperties.solidFill = COLOR_CATEGORIA["Banco tradicional"]
    for i in range(top):
        pt = DataPoint(idx=i)
        pt.graphicalProperties.solidFill = COLOR_CATEGORIA[sim.iloc[i]["Categoria"]]
        s.dPt.append(pt)
    linea = LineChart()
    linea.add_data(Reference(ws_sim, min_col=15, min_row=4, max_row=4 + top), titles_from_data=True)
    linea.series[0].graphicalProperties.line.solidFill = "DC2626"
    linea.series[0].graphicalProperties.line.dashStyle = "dash"
    linea.series[0].graphicalProperties.line.width = 19000
    linea.series[0].smooth = False
    bar += linea
    bar.legend.position = "b"
    bar.height, bar.width = 10, 27
    ws_res["B34"] = "RENTABILIDAD NETA EA POR ENTIDAD · TOP 15 VS PROMEDIO NETO DEL SISTEMA"
    ws_res["B34"].font = Font(bold=True, size=8, color=xp.GRIS_TENUE)
    ws_res.add_chart(bar, "B35")
    xp.nota(ws_res, "B57", "Colores: gris oscuro = banco tradicional, verde azulado = banco digital, ámbar = fintech. "
            "La línea roja punteada es el promedio del sistema (BanRep) con el mismo cálculo neto. " + ADVERTENCIA, 12, 40)

    # ---- Formato institucional ----
    for hoja, rol in ((ws_res, "salida"), (ws_ctl, "control"), (ws_par, "entrada"), (ws_tas, "dato"), (ws_sim, "calculo"),
                      (ws_flu, "calculo"), (ws_aho, "calculo"), (ws_ref, "dato"), (ws_fue, "dato")):
        xp.formato_hoja(hoja, rol)
    xp.semaforo(ws_flu, "C11:C12")
    xp.estandarizar_fuentes(wb)
    wb.save(ruta)


# ----------------------------------------------------------------------------------------------------------------
# Markdown
# ----------------------------------------------------------------------------------------------------------------
def nombre_visible(fila) -> str:
    marca = fila["MarcaComercial"] if "MarcaComercial" in fila else ""
    if marca and marca.lower() not in fila["Entidad"].lower():
        return f"{fila['Entidad']} ({marca})"
    return fila["Entidad"]


def generar_md(d: dict, ruta_md: Path, ruta_xlsx: Path, checks: list[dict]) -> None:
    sim, ah, mej = d["sim"], d["ah"], d["mejor"]
    plantilla = (SCRIPT_DIR / "plantilla_md.md").read_text(encoding="utf-8")
    monto, plazo = d["monto"], d["plazo"]
    prom_sis = d["sistema"]["valor"] / 100

    top10 = tabla_md(
        ["#", "Entidad", "Categoría", "Tasa EA", "Interés neto", "Rentabilidad real EA"],
        [[int(r["Ranking"]), nombre_visible(r), r["Categoria"], f_pct(r["TasaEA"]), f_cop(r["InteresNeto"]),
          f_pct(r["RentRealEA"])] for _, r in sim.head(10).iterrows()])

    filas_g = []
    for cat in CATEGORIAS:
        g = sim[sim["Categoria"] == cat]
        if g.empty:
            filas_g.append([cat, 0, "Sin dato", "Sin dato", "Sin dato", "Sin dato"])
            continue
        m = g.iloc[0]
        filas_g.append([cat, len(g), f_pct(g["TasaEA"].mean()), f_pct(g["RentRealEA"].mean()),
                        nombre_visible(m), f_pct(m["TasaEA"])])
    grupos = tabla_md(["Categoría", "Entidades con dato", "Tasa EA promedio", "Rentabilidad real EA promedio",
                       "Mejor entidad", "Su tasa EA"], filas_g)

    bloque_inflacion = "\n".join([
        f"- Inflación a 12 meses (DANE, {f_fecha(d['fecha_ipc'])}): {f_pct(d['ipc12m'])}."
        + (f" {d['ipc_nota']}" if d["ipc_nota"] else ""),
        f"- Promedio del sistema a {d['sistema']['nombre'].split(' a ')[-1].split(',')[0]} (BanRep, {f_fecha(d['sistema']['fecha'])}): {f_pct(prom_sis)} EA bruta.",
        f"- Mejor entidad: {nombre_visible(mej)} con {f_pct(mej['TasaEA'])} EA bruta, {f_pp(mej['DifVsSistema'])} frente al sistema.",
        f"- Rentabilidad neta EA de la mejor (tras retención de {f_pct(d['retencion'], 0)}): {f_pct(mej['RentNetaEA'])}; "
        f"real (descontada la inflación): {f_pct(mej['RentRealEA'])}.",
        f"- Tasa de política monetaria: {f_pct(d['tpm']['valor'] / 100, 2)}; IBR a 3 meses: {f_pct(d['ibr']['valor'] / 100, 2)}.",
    ])

    sel = ah[ah["Entidad"].isin(ENTIDADES_AHORRO_MD)]
    filas_a = [[r["Entidad"], f_pct(r["TasaEA"]), f_cop(r["InteresNetoAhorro"]), f_cop(r["DifVsMejorCDT"])]
               for _, r in sel.iterrows()]
    tabla_ah = tabla_md(["Cuenta de ahorro", "Tasa EA", "Interés neto en el plazo", "Diferencia vs mejor CDT"], filas_a)
    tabla_ah += f"\n\nMejor CDT ({nombre_visible(mej)}): interés neto {f_cop(mej['InteresNeto'])}."

    alertas = []
    reza = [f"{r['Entidad']} ({int(r['RezagoCortes'])} corte de rezago)" for _, r in d["tasas"].iterrows()
            if r["Estado"] == "OK" and pd.notna(r["RezagoCortes"]) and r["RezagoCortes"] > 0]
    if reza:
        alertas.append("- Datos con rezago (se usó el último corte disponible): " + ", ".join(reza) + ".")
    baja = [r["Entidad"] for _, r in d["tasas"].iterrows() if r["Alerta"] == "Baja representatividad"]
    if baja:
        alertas.append("- Baja representatividad (menos de $50 millones captados ese día): " + ", ".join(baja) + ".")
    exc = sim[sim["AlertaFogafin"] != ""]
    if not exc.empty:
        alertas.append(f"- El valor final ({f_cop(mej['ValorFinal'])}) supera la cobertura del seguro de depósitos "
                       f"({f_cop(COBERTURA_FOGAFIN)} por persona y entidad, capital más intereses): la parte excedente no "
                       "está asegurada. Considere repartir el monto entre entidades.")
    else:
        alertas.append(f"- El valor final ({f_cop(mej['ValorFinal'])}) está dentro de la cobertura del seguro de depósitos "
                       f"({f_cop(COBERTURA_FOGAFIN)} por persona y entidad).")
    sd = d["tasas"][d["tasas"]["Estado"] != "OK"]
    if not sd.empty:
        alertas.append("- Sin dato de CDT en este plazo: " + ", ".join(f"{r.Entidad} ({r.Estado})" for r in sd.itertuples()) + ".")
    if d["faltantes"]:
        alertas.append("- Entidades de entidades.csv que ya no aparecen en el dataset: " + ", ".join(d["faltantes"]) + ".")
    for a in ADVERTENCIAS:
        alertas.append(f"- {a}")
    if sfc.plazo_es_aproximado(plazo):
        alertas.append(f"- El plazo de {plazo} días no coincide con un plazo publicado; se usó el rango más cercano ({sfc.DESC_SUBCUENTA[d['subcuenta']]}).")
    fallos = [c for c in checks if c["Resultado"] != "OK"]
    for c in fallos:
        alertas.append(f"- Validación {c['Resultado'].lower()}: {c['Validacion']} ({c['Detalle']}).")

    fuentes_md = "\n".join(f"- {f['Fuente']}: {f['URL']} (consulta {f['FechaHoraConsulta']})" for f in d["fuentes"])

    resultado = (f"Un CDT de {f_cop(monto)} a {plazo} días rinde más hoy en {nombre_visible(mej)}: {f_pct(mej['TasaEA'])} EA, "
                 f"con un interés neto de {f_cop(mej['InteresNeto'])} tras retención (rentabilidad real {f_pct(mej['RentRealEA'])} EA).")
    valores = {
        "RESULTADO_UNA_LINEA": resultado, "MONTO": f_cop(monto), "PLAZO": str(plazo),
        "PERIODICIDAD": ETIQUETA_PERIODICIDAD[d["periodicidad"]], "FECHA_CORTE": f_fecha(d["cortes"][0]),
        "TABLA_TOP10": top10, "TABLA_GRUPOS": grupos, "BLOQUE_INFLACION": bloque_inflacion, "TABLA_AHORRO": tabla_ah,
        "ALERTAS": "\n".join(alertas), "ADVERTENCIA": ADVERTENCIA, "FUENTES": fuentes_md,
        "EXCEL": ruta_xlsx.name, "FECHA_CONSULTA": d["ahora"].strftime("%Y-%m-%d %H:%M"),
        "SUBCUENTA_DESC": sfc.DESC_SUBCUENTA.get(d["subcuenta"], ""),
    }
    for k, v in valores.items():
        plantilla = plantilla.replace("{{" + k + "}}", v)
    ruta_md.write_text(plantilla, encoding="utf-8")


# ----------------------------------------------------------------------------------------------------------------
# Relleno y validación de las interpretaciones escritas por el modelo
# ----------------------------------------------------------------------------------------------------------------
# ----------------------------------------------------------------------------------------------------------------
# Validación con Excel (opcional, solo Windows con Excel instalado)
# ----------------------------------------------------------------------------------------------------------------
def validar_con_excel(ruta_xlsx: Path, d: dict, guardar: bool = True) -> dict:
    ps1 = SCRIPT_DIR / "validar_excel.ps1"
    if sys.platform != "win32" or not ps1.exists():
        return {"disponible": False, "motivo": "Solo Windows con Excel instalado"}
    with tempfile.TemporaryDirectory() as tmp:
        salida = Path(tmp) / "val.json"
        cmd = ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(ps1), "-Ruta", str(ruta_xlsx),
               "-Salida", str(salida)] + (["-Guardar"] if guardar else [])
        try:
            subprocess.run(cmd, capture_output=True, text=True, timeout=180)
            r = json.loads(salida.read_text(encoding="utf-8-sig"))
        except Exception as e:  # noqa: BLE001
            return {"disponible": False, "motivo": f"No se pudo validar con Excel: {e}"}
    if not r.get("disponible"):
        return r
    sim, mej = d["sim"], d["mejor"]
    fallos = []
    if r["MejorEntidad"] != mej["Entidad"]:
        fallos.append(f"MejorEntidad Excel={r['MejorEntidad']} Python={mej['Entidad']}")
    for k, py in (("TasaMejor", mej["TasaEA"]), ("MejorInteresNeto", mej["InteresNeto"]),
                  ("SumaInteresNeto", sim["InteresNeto"].sum()), ("SumaRentNetaEA", sim["RentNetaEA"].sum()),
                  ("RentRealMejor", mej["RentRealEA"])):
        if abs(r[k] - py) > max(0.01, abs(py) * 1e-9):
            fallos.append(f"{k} Excel={r[k]} Python={py}")
    if abs(r["SumXfd"]) >= 1 or r["ValidaXfd"] != "OK":
        fallos.append(f"ΣXfd Excel={r['SumXfd']} ({r['ValidaXfd']})")
    if r["Errores"] != 0:
        fallos.append(f"{r['Errores']} celdas con error en Excel")
    if abs(r["InteresNetoMonto2x"] / r["MejorInteresNeto"] - 2) > 1e-9:
        fallos.append("Al duplicar el Monto el interés neto no se duplicó")
    r["fallos"] = fallos
    return r


# ----------------------------------------------------------------------------------------------------------------
# Programa principal
# ----------------------------------------------------------------------------------------------------------------
def ejecutar(a: argparse.Namespace) -> int:
    if a.plazo < 30:
        print("ERROR: el plazo mínimo es 30 días (los CDT en Colombia se emiten desde 30 días).")
        return 2
    dp = dias_periodo_de(a.periodicidad, a.plazo)
    if a.plazo % dp != 0:
        print(f"ERROR: con periodicidad '{a.periodicidad}' (cada {dp} días) el plazo debe ser múltiplo de {dp}. "
              f"Pruebe plazo 360, o periodicidad 'vencimiento'. Plazo pedido: {a.plazo}.")
        return 2
    if a.plazo // dp > MAX_PERIODOS_FLUJO:
        print(f"ERROR: máximo {MAX_PERIODOS_FLUJO} periodos de pago; reduzca el plazo o use otra periodicidad.")
        return 2
    try:
        d = recolectar(a.monto, a.plazo, a.periodicidad, a.retencion, a.base_dias)
        d = calcular_todo(d)
    except FuenteError as e:
        print("ERROR_FUENTE:", e)
        return 3
    salida = Path(a.salida) if a.salida else SALIDA_DEFECTO
    salida.mkdir(parents=True, exist_ok=True)
    base_nombre = f"CDTLive_{a.monto}_{a.plazo}d_{date.today().isoformat()}"
    ruta_xlsx, aviso_archivo = xp.ruta_libre(salida / f"{base_nombre}.xlsx")
    ruta_md = ruta_xlsx.with_suffix(".md")
    if aviso_archivo:
        ADVERTENCIAS.append(aviso_archivo)
    checks = validaciones_python(d)
    construir_excel(d, ruta_xlsx, checks)
    val = {"disponible": False, "motivo": "omitida (--sin-validar-excel)"}
    if not a.sin_validar_excel:
        val = validar_con_excel(ruta_xlsx, d, guardar=True)
    if val.get("disponible"):
        estado = "OK" if not val["fallos"] else "ERROR"
        checks.append({"Validacion": "Excel recalculado coincide con Python; al duplicar el Monto todo se recalcula",
                       "Resultado": estado, "Detalle": "; ".join(val["fallos"]) or
                       f"{val['Errores']} errores de fórmula; interés neto mejor {val['MejorInteresNeto']:.2f}"})
    else:
        checks.append({"Validacion": "Excel recalculado coincide con Python", "Resultado": "OMITIDA",
                       "Detalle": val.get("motivo", "")})
    if val.get("disponible"):
        d["estado_conciliacion"] = "OK" if not val["fallos"] else "ERROR: " + "; ".join(val["fallos"])[:200]
    else:
        d["estado_conciliacion"] = "OMITIDA: " + val.get("motivo", "")
    construir_excel(d, ruta_xlsx, checks)  # reescribe con el resultado de la validación en Controles y Fuentes
    if val.get("disponible"):
        val2 = validar_con_excel(ruta_xlsx, d, guardar=True)  # deja valores calculados guardados en el archivo
        val["fallos"] = val["fallos"] or val2.get("fallos", [])
    generar_md(d, ruta_md, ruta_xlsx, checks)

    m = d["mejor"]
    print("RESUMEN_OK")
    print(f"excel={ruta_xlsx}")
    print(f"md={ruta_md}")
    print(f"fecha_corte={d['cortes'][0]}")
    print(f"resultado={f_cop(a.monto)} a {a.plazo} dias: mejor {nombre_visible(m)}, {f_pct(m['TasaEA'])} EA, interes neto {f_cop(m['InteresNeto'])}")
    print(f"rentabilidad_real_ea={f_pct(m['RentRealEA'])}; promedio_sistema={f_pct(d['sistema']['valor'] / 100)}; ipc12m={f_pct(d['ipc12m'])}")
    print(f"entidades_con_dato={len(d['sim'])}; sin_dato={', '.join(d['tasas'][d['tasas']['Estado'] != 'OK']['Entidad'])}")
    for c in checks:
        print(f"VALIDACION [{c['Resultado']}] {c['Validacion']}: {c['Detalle']}")
    if aviso_archivo:
        print("AVISO_ARCHIVO:", aviso_archivo)
    print("SIGUIENTE_PASO: escriba las 4 interpretaciones y ejecute --rellenar (ver SKILL.md, paso 4)")
    return 0 if all(c["Resultado"] in ("OK", "OMITIDA", "ADVERTENCIA") for c in checks) else 4


def main() -> int:
    p = argparse.ArgumentParser(description="CDTLive: comparador de CDT con datos oficiales de Colombia")
    p.add_argument("--monto", type=int, default=10_000_000, help="Capital en COP (defecto 10.000.000)")
    p.add_argument("--plazo", type=int, default=360, help="Plazo en días (defecto 360; mínimo 30)")
    p.add_argument("--periodicidad", choices=list(PERIODICIDADES), default="vencimiento")
    p.add_argument("--retencion", type=float, default=0.04, help="Retención en la fuente sobre intereses (defecto 0.04)")
    p.add_argument("--base-dias", type=int, default=365)
    p.add_argument("--salida", help="Carpeta de salida (defecto: Planeación Financiera/salidas/CDTLive)")
    p.add_argument("--sin-validar-excel", action="store_true", help="Omite la validación con Excel (COM)")
    p.add_argument("--rellenar", metavar="RUTA_MD", help="Inserta las interpretaciones en el informe")
    for k in ("grupos", "inflacion", "ahorro", "alertas"):
        p.add_argument(f"--{k}", help=f"Interpretación de la sección {k} (máx. 3 frases, solo cifras del informe)")
    a = p.parse_args()
    if a.rellenar:
        textos = {k.upper(): getattr(a, k) for k in ("grupos", "inflacion", "ahorro", "alertas") if getattr(a, k) is not None}
        if not textos:
            print("ERROR: indique al menos una de --grupos --inflacion --ahorro --alertas")
            return 2
        return rellenar(Path(a.rellenar), textos)
    return ejecutar(a)


if __name__ == "__main__":
    sys.exit(main())
